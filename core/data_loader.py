# -*- coding: utf-8 -*-
"""
数据加载与预处理模块
"""

try:
    import pandas as pd
except ModuleNotFoundError:
    pd = None

from openpyxl import load_workbook
from core.config import (
    DEFAULT_DATA_FILE,
    SCALE_FACTOR,
    MIN_CUT_GAP,
    VERTICAL_CUT_INCLUSIVE,
)


class Item:
    """单个小板的数据结构"""
    def __init__(self, item_id, type_id, width, length, weight):
        self.item_id = item_id      # 唯一ID
        self.type_id = type_id      # 产品类型ID
        self.width = width          # 宽度
        self.length = length        # 长度
        self.weight = weight        # 重量
    
    def __repr__(self):
        return f"Item(type={self.type_id}, w={self.width}, l={self.length})"


class _SimpleColumn:
    def __init__(self, values):
        self.iloc = values


class SimpleTypeInfo:
    """Lightweight fallback used when pandas is unavailable."""

    def __init__(self, demand_info):
        self._data = {key: list(values) for key, values in demand_info.items()}

    def __getitem__(self, key):
        return _SimpleColumn(self._data[key])

    def __len__(self):
        return len(self._data.get("Width", []))


def load_demand_from_excel(filepath='产品数据.xlsx', sheet_num=1):
    """
    从Excel文件加载需求数据
    
    返回:
        demand_info: dict，包含Width, Length, num, Weight列表
    """
    if filepath is None:
        filepath = DEFAULT_DATA_FILE
    wb = load_workbook(filepath, data_only=True)
    ws = wb[f'Sheet{sheet_num}']
    
    item_num = ws['B1'].value
    if item_num is None or item_num <= 0:
        raise ValueError("产品种类数量输入不正确")
    
    width_list = []
    length_list = []
    num_list = []
    weight_list = []
    
    for i in range(item_num):
        width = ws[f'A{i+3}'].value
        length = ws[f'B{i+3}'].value
        num = ws[f'C{i+3}'].value
        weight = ws[f'D{i+3}'].value
        if width is not None:
            if VERTICAL_CUT_INCLUSIVE and width < MIN_CUT_GAP:
                raise ValueError(
                    f"Type {i + 1} width {width}mm is smaller than the minimum vertical cut spacing {MIN_CUT_GAP}mm"
                )
            if (not VERTICAL_CUT_INCLUSIVE) and width <= MIN_CUT_GAP:
                raise ValueError(
                    f"Type {i + 1} width {width}mm must be greater than {MIN_CUT_GAP}mm"
                )
        
        if width is None or width <= 0:
            raise ValueError(f"第{i+1}种产品宽度输入不正确")
        if length is None or length <= 0:
            raise ValueError(f"第{i+1}种产品长度输入不正确")
        if num is None or num <= 0:
            raise ValueError(f"第{i+1}种产品数量输入不正确")
        if weight is None or weight <= 0:
            raise ValueError(f"第{i+1}种产品重量输入不正确")
        
        width_list.append(width)
        length_list.append(length)
        num_list.append(num)
        weight_list.append(weight)
    
    demand_info = {
        'Width': width_list,
        'Length': length_list,
        'num': num_list,
        'Weight': weight_list
    }
    
    return demand_info


def expand_demand(demand_info):
    """
    将需求展开成单个小板的列表（基因编码）
    
    输入:
        demand_info: dict，包含Width, Length, num, Weight
    
    返回:
        items: list of Item，所有待切割的小板列表
        type_info: DataFrame，原始类型信息
    """
    items = []
    item_id = 0
    
    if pd is not None:
        type_info = pd.DataFrame(demand_info)
    else:
        type_info = SimpleTypeInfo(demand_info)
    
    for type_id in range(len(demand_info['Width'])):
        width = demand_info['Width'][type_id]
        length = demand_info['Length'][type_id]
        weight = demand_info['Weight'][type_id]
        num = demand_info['num'][type_id]
        if VERTICAL_CUT_INCLUSIVE and width < MIN_CUT_GAP:
            raise ValueError(
                f"Type {type_id + 1} width {width}mm is smaller than the minimum vertical cut spacing {MIN_CUT_GAP}mm"
            )
        if (not VERTICAL_CUT_INCLUSIVE) and width <= MIN_CUT_GAP:
            raise ValueError(
                f"Type {type_id + 1} width {width}mm must be greater than {MIN_CUT_GAP}mm"
            )
        
        # 按比例缩放数量
        scaled_num = max(1, int(num / SCALE_FACTOR))
        
        for _ in range(scaled_num):
            item = Item(item_id, type_id, width, length, weight)
            items.append(item)
            item_id += 1
    
    return items, type_info


def get_demand_interactive():
    """交互式获取需求数据"""
    sheet_num = int(input("请输入产品数据表: "))
    demand_info = load_demand_from_excel(sheet_num=sheet_num)
    return demand_info


if __name__ == "__main__":
    # 测试数据加载
    demand = load_demand_from_excel(sheet_num=1)
    print("原始需求:")
    if pd is not None:
        print(pd.DataFrame(demand))
    else:
        print(demand)
    
    items, type_info = expand_demand(demand)
    print(f"\n展开后的小板数量: {len(items)}")
    print(f"前10个小板: {items[:10]}")
