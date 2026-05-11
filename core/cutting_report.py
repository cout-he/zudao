from collections import defaultdict
from pathlib import Path
import unicodedata

from core.config import PANEL_WIDTH, SCALE_FACTOR
from core.decoder import merge_same_pattern_strips


STEEL_DENSITY_T_PER_MM3 = 7.85e-9


def _collect_lane_groups(strip):
    groups = []
    group_index = {}
    for item in strip.items:
        key = (item.type_id, item.width, item.length)
        if key not in group_index:
            group_index[key] = len(groups)
            groups.append(
                {
                    "type_id": item.type_id,
                    "width": item.width,
                    "length": item.length,
                    "lane_count": 0,
                }
            )
        groups[group_index[key]]["lane_count"] += 1
    return groups


def _format_mm(value):
    if float(value).is_integer():
        return str(int(value))
    return f"{float(value):.2f}".rstrip("0").rstrip(".")


def _format_m(value_mm):
    return f"{float(value_mm) / 1000.0:.3f}"


def _format_ton(value_ton):
    return f"{float(value_ton):.6f}"


def _format_width_combo(groups):
    return " + ".join(
        [str(int(group["width"])) for group in groups for _ in range(group["lane_count"])]
    )


def _display_width(value):
    width = 0
    for char in str(value):
        width += 2 if unicodedata.east_asian_width(char) in {"F", "W"} else 1
    return width


def _pad_display(value, width, align="left"):
    text = str(value)
    padding = max(int(width) - _display_width(text), 0)
    if align == "right":
        return " " * padding + text
    return text + " " * padding


def _format_columns(values, widths, aligns=None, gap="  "):
    aligns = aligns or ["left"] * len(values)
    return gap.join(
        _pad_display(value, width, align)
        for value, width, align in zip(values, widths, aligns)
    )


def _format_box_table(headers, rows, aligns=None):
    aligns = aligns or ["left"] * len(headers)
    table_rows = [headers, *rows]
    widths = [
        max(_display_width(row[col]) for row in table_rows)
        for col in range(len(headers))
    ]

    def border():
        return "+" + "+".join("-" * (width + 2) for width in widths) + "+"

    def line(values):
        cells = [
            " " + _pad_display(value, width, align) + " "
            for value, width, align in zip(values, widths, aligns)
        ]
        return "|" + "|".join(cells) + "|"

    lines = [border(), line(headers), border()]
    lines.extend(line(row) for row in rows)
    lines.append(border())
    return lines


def _format_knife_positions(groups):
    positions = []
    cursor = 0
    for group in groups:
        for _ in range(group["lane_count"]):
            cursor += int(group["width"])
            positions.append(cursor)
    return " / ".join(str(pos) for pos in positions)


def _format_table(rows, key_header="项目", value_header="内容"):
    table_rows = [(key_header, value_header), *rows]
    key_width = max(_display_width(row[0]) for row in table_rows)
    lines = [_format_columns([key, value], [key_width, 0]) for key, value in table_rows]
    return lines


def _parse_thickness_from_demand(demand):
    specs = demand.get("Spec") or []
    for spec in specs:
        text = str(spec).replace("×", "*").replace("x", "*").replace("X", "*")
        parts = [part for part in text.split("*") if part.strip()]
        if parts:
            try:
                return float(parts[0])
            except ValueError:
                continue
    return None


def _infer_weight_unit(demand):
    weights = [float(value) for value in demand.get("Weight", []) if value is not None]
    if not weights:
        return "kg"

    thickness = _parse_thickness_from_demand(demand)
    if not thickness:
        return "kg"

    theoretical_kg = 0.0
    actual_weight = 0.0
    for width, length, count, weight in zip(
        demand.get("Width", []),
        demand.get("Length", []),
        demand.get("num", []),
        demand.get("Weight", []),
    ):
        if weight is None:
            continue
        theoretical_kg += (
            float(width)
            * float(length)
            * thickness
            * STEEL_DENSITY_T_PER_MM3
            * 1000.0
            * int(count)
        )
        actual_weight += float(weight)

    if theoretical_kg <= 0 or actual_weight <= 0:
        return "kg"

    kg_ratio = actual_weight / theoretical_kg
    ton_ratio = actual_weight * 1000.0 / theoretical_kg
    return "kg" if abs(kg_ratio - 1.0) <= abs(ton_ratio - 1.0) else "t"


def _input_weight_to_tons(value, weight_unit):
    if value is None:
        return 0.0
    value = float(value)
    return value if weight_unit == "t" else value / 1000.0


def _estimate_consumed_weight_tons(width_mm, length_mm, demand, weight_unit):
    thickness = _parse_thickness_from_demand(demand)
    if thickness:
        return float(width_mm) * float(length_mm) * float(thickness) * STEEL_DENSITY_T_PER_MM3

    order_weight_tons = sum(
        _input_weight_to_tons(value, weight_unit) for value in demand.get("Weight", [])
    )
    demand_area = sum(
        float(width) * float(length) * int(count)
        for width, length, count in zip(
            demand.get("Width", []),
            demand.get("Length", []),
            demand.get("num", []),
        )
    )
    if demand_area <= 0:
        return 0.0
    return order_weight_tons * (float(width_mm) * float(length_mm)) / demand_area


def _piece_weight_tons(demand, type_id, weight_unit):
    weights = demand.get("Weight") or []
    counts = demand.get("num") or []
    if type_id >= len(weights) or type_id >= len(counts) or int(counts[type_id]) <= 0:
        return 0.0
    return _input_weight_to_tons(weights[type_id], weight_unit) / int(counts[type_id])


def _build_stage_sketch(groups):
    cells = [f"T{group['type_id'] + 1} {int(group['width'])}" for group in groups for _ in range(group["lane_count"])]
    used_width = sum(group["width"] * group["lane_count"] for group in groups)
    waste_width = PANEL_WIDTH - used_width
    if waste_width > 0:
        cells.append(f"余宽 {int(waste_width)}")
    return " | ".join(cells)


def calculate_produced_totals(decoder_mode, solution):
    produced_totals = defaultdict(int)

    if decoder_mode == "stage_based":
        for strip in solution["strips"]:
            actual_stage_length = strip.strip_length * SCALE_FACTOR
            groups = _collect_lane_groups(strip)
            for group in groups:
                pieces_per_lane = actual_stage_length // group["length"]
                lane_total = pieces_per_lane * group["lane_count"]
                produced_totals[group["type_id"]] += int(lane_total)
    else:
        merged_strips, repeat_counts = merge_same_pattern_strips(solution["strips"])
        for strip, repeat_count in zip(merged_strips, repeat_counts):
            actual_repeat = repeat_count * SCALE_FACTOR
            groups = _collect_lane_groups(strip)
            for group in groups:
                lane_total = actual_repeat * group["lane_count"]
                produced_totals[group["type_id"]] += int(lane_total)

    return dict(produced_totals)


def summarize_production(demand, solution, decoder_mode):
    produced_totals = calculate_produced_totals(decoder_mode, solution)
    order_total = sum(int(num) for num in demand["num"])
    produced_total = sum(int(value) for value in produced_totals.values())
    shortage_total = 0
    over_total = 0
    demand_area = 0.0
    produced_area = 0.0

    for idx, (width, length, demand_num) in enumerate(
        zip(demand["Width"], demand["Length"], demand["num"])
    ):
        ordered = int(demand_num)
        produced = int(produced_totals.get(idx, 0))
        shortage_total += max(ordered - produced, 0)
        over_total += max(produced - ordered, 0)
        demand_area += float(width) * float(length) * ordered
        produced_area += float(width) * float(length) * produced

    real_total_length = float(solution["total_length"]) * SCALE_FACTOR
    panel_area = PANEL_WIDTH * real_total_length
    demand_utilization = 100 * demand_area / panel_area if panel_area else 0.0
    actual_utilization = 100 * produced_area / panel_area if panel_area else 0.0
    weight_unit = _infer_weight_unit(demand)
    input_weight_total_tons = sum(
        _input_weight_to_tons(value, weight_unit) for value in demand.get("Weight", [])
    )
    consumed_weight_tons = _estimate_consumed_weight_tons(
        PANEL_WIDTH,
        real_total_length,
        demand,
        weight_unit,
    )

    return {
        "produced_totals": produced_totals,
        "order_total": int(order_total),
        "produced_total": int(produced_total),
        "shortage_total": int(shortage_total),
        "over_total": int(over_total),
        "demand_area": float(demand_area),
        "produced_area": float(produced_area),
        "real_total_length": float(real_total_length),
        "panel_area": float(panel_area),
        "demand_utilization": float(demand_utilization),
        "actual_utilization": float(actual_utilization),
        "weight_unit": weight_unit,
        "input_weight_total_tons": float(input_weight_total_tons),
        "consumed_weight_tons": float(consumed_weight_tons),
    }


def _format_task_section(sheet_num, decoder_mode, demand, solution, production_summary):
    thickness = _parse_thickness_from_demand(demand)
    stage_count = len(solution["strips"]) if decoder_mode == "stage_based" else int(solution["num_strips"])
    product_name = str(demand.get("Name") or "-")
    rows = [
        ("任务编号", sheet_num),
        ("品名", product_name),
        ("母卷规格", f"{_format_mm(thickness) if thickness else '-'} x {PANEL_WIDTH} mm"),
        ("输入重量单位", "kg，报告已换算为吨" if production_summary["weight_unit"] == "kg" else "吨"),
        ("订单重量合计", f"{_format_ton(production_summary['input_weight_total_tons'])} 吨"),
        ("总走料长度", f"{_format_m(production_summary['real_total_length'])} m"),
        ("预计用料重量", f"{_format_ton(production_summary['consumed_weight_tons'])} 吨"),
        ("面积利用率", f"{production_summary['actual_utilization']:.2f}%"),
        ("生产阶段数", stage_count),
    ]
    return ["1. 生产任务信息", *_format_table(rows), ""]


def _format_order_check_section(demand, produced_totals):
    lines = ["2. 订单与产出核对"]
    aligns = ["left", "left", "right", "right", "right", "left"]
    rows = []

    for idx, (width, length, demand_num) in enumerate(
        zip(demand["Width"], demand["Length"], demand["num"]),
        start=1,
    ):
        produced = int(produced_totals.get(idx - 1, 0))
        gap = produced - int(demand_num)
        if gap == 0:
            note = "与订单一致"
            gap_text = "0"
        elif gap > 0:
            note = f"多 {gap} 件，作为余量"
            gap_text = f"+{gap}"
        else:
            note = f"需补切 {abs(gap)} 件"
            gap_text = str(gap)
        spec = f"{_format_mm(width)} x {_format_mm(length)}"
        rows.append([f"T{idx}", spec, int(demand_num), produced, gap_text, note])

    lines.extend(
        _format_box_table(
            ["产品编号", "规格 mm", "订单数量", "本单产出", "差异", "处理说明"],
            rows,
            aligns,
        )
    )
    lines.append("")
    return lines


def _format_stage_output_rows(rows):
    aligns = ["left", "right", "right", "right", "right", "left"]
    headers = ["产品编号", "每道产出", "本阶段产出", "产出重量(吨)", "单道尾料(mm)", "说明"]
    return _format_box_table(headers, rows, aligns)


def _build_compact_mode_section(demand, strips, repeat_counts):
    lines = ["3. 排刀执行说明"]
    produced_totals = defaultdict(int)
    weight_unit = _infer_weight_unit(demand)

    for idx, (strip, repeat_count) in enumerate(zip(strips, repeat_counts), start=1):
        actual_repeat = repeat_count * SCALE_FACTOR
        actual_length = float(strip.strip_length) * float(actual_repeat)
        groups = _collect_lane_groups(strip)
        used_width = sum(group["width"] * group["lane_count"] for group in groups)
        waste_width = PANEL_WIDTH - used_width
        stage_weight = _estimate_consumed_weight_tons(PANEL_WIDTH, actual_length, demand, weight_unit)

        lines.append(f"阶段 {idx}：连续走料模式")
        rows = [
            ("排刀组合", _format_width_combo(groups)),
            ("刀数", f"{sum(group['lane_count'] for group in groups)} 道"),
            ("占用宽度", f"{int(used_width)} mm"),
            ("边部余宽", f"{int(waste_width)} mm"),
            ("单条长度", f"{_format_mm(strip.strip_length)} mm"),
            ("连续条数", f"{int(actual_repeat)} 条"),
            ("走料长度", f"{_format_m(actual_length)} m"),
            ("预计用料重量", f"{_format_ton(stage_weight)} 吨"),
            ("刀位累计位置", f"{_format_knife_positions(groups)} mm"),
        ]
        lines.extend(_format_table(rows))

        output_rows = []
        for group in groups:
            lane_total = int(actual_repeat * group["lane_count"])
            produced_totals[group["type_id"]] += lane_total
            product_weight = lane_total * _piece_weight_tons(demand, group["type_id"], weight_unit)
            tail_length = max(0, float(strip.strip_length) - float(group["length"]))
            output_rows.append(
                [
                    f"T{group['type_id'] + 1}",
                    int(actual_repeat),
                    lane_total,
                    _format_ton(product_weight),
                    _format_mm(tail_length),
                    f"{group['lane_count']} 道并切",
                ]
            )
        lines.extend(_format_stage_output_rows(output_rows))
        lines.append(f"示意图：母卷宽度 {PANEL_WIDTH} mm | {_build_stage_sketch(groups)}")
        lines.append("")

    return lines, produced_totals


def _build_stage_mode_section(demand, strips):
    lines = ["3. 排刀执行说明"]
    produced_totals = defaultdict(int)
    weight_unit = _infer_weight_unit(demand)

    for idx, strip in enumerate(strips, start=1):
        actual_stage_length = strip.strip_length * SCALE_FACTOR
        groups = _collect_lane_groups(strip)
        used_width = sum(group["width"] * group["lane_count"] for group in groups)
        waste_width = PANEL_WIDTH - used_width
        stage_weight = _estimate_consumed_weight_tons(
            PANEL_WIDTH,
            actual_stage_length,
            demand,
            weight_unit,
        )
        product_labels = sorted({f"T{group['type_id'] + 1}" for group in groups})

        lines.append(f"阶段 {idx}：{' + '.join(product_labels)} 同一组刀位连续走料")
        rows = [
            ("排刀组合", _format_width_combo(groups)),
            ("刀数", f"{sum(group['lane_count'] for group in groups)} 道"),
            ("占用宽度", f"{int(used_width)} mm"),
            ("边部余宽", f"{int(waste_width)} mm"),
            ("定尺长度", " / ".join(f"T{group['type_id'] + 1}:{_format_mm(group['length'])} mm" for group in groups)),
            ("走料长度", f"{_format_m(actual_stage_length)} m"),
            ("预计用料重量", f"{_format_ton(stage_weight)} 吨"),
            ("刀位累计位置", f"{_format_knife_positions(groups)} mm"),
        ]
        lines.extend(_format_table(rows))

        output_rows = []
        for group in groups:
            pieces_per_lane = int(actual_stage_length // group["length"])
            lane_total = int(pieces_per_lane * group["lane_count"])
            tail_length = actual_stage_length - pieces_per_lane * group["length"]
            produced_totals[group["type_id"]] += lane_total
            product_weight = lane_total * _piece_weight_tons(demand, group["type_id"], weight_unit)
            output_rows.append(
                [
                    f"T{group['type_id'] + 1}",
                    pieces_per_lane,
                    lane_total,
                    _format_ton(product_weight),
                    _format_mm(tail_length),
                    f"{group['lane_count']} 道并切",
                ]
            )
        lines.extend(_format_stage_output_rows(output_rows))
        lines.append(f"示意图：母卷宽度 {PANEL_WIDTH} mm | {_build_stage_sketch(groups)}")
        lines.append("")

    return lines, produced_totals


def _format_execution_notes(demand, produced_totals):
    shortage_notes = []
    over_notes = []
    for idx, demand_num in enumerate(demand["num"], start=1):
        produced = int(produced_totals.get(idx - 1, 0))
        gap = produced - int(demand_num)
        if gap < 0:
            shortage_notes.append(f"T{idx} 少 {abs(gap)} 件，需要单独补切")
        elif gap > 0:
            over_notes.append(f"T{idx} 多 {gap} 件，可作为余量件")

    rows = [
        ("换刀顺序", "按阶段编号顺序执行；每阶段内同一组刀位连续走料"),
        ("补切要求", "；".join(shortage_notes) if shortage_notes else "无"),
        ("多产说明", "；".join(over_notes) if over_notes else "无"),
        ("注意事项", "阶段执行过程中不要随意改变刀位组合"),
    ]
    return ["4. 换刀与补切提示", *_format_table(rows), ""]


def _format_confirm_section():
    rows = [
        ("排刀确认", ""),
        ("生产确认", ""),
        ("数量核对", ""),
        ("补切确认", ""),
        ("日期", ""),
    ]
    return ["5. 车间确认栏", *_format_table(rows), ""]


def _make_border():
    from openpyxl.styles import Border, Side

    side = Side(style="thin", color="B7B7B7")
    return Border(left=side, right=side, top=side, bottom=side)


def _style_range(ws, cell_range, *, fill=None, bold=False, align="left"):
    from openpyxl.styles import Alignment, Font, PatternFill

    border = _make_border()
    for row in ws[cell_range]:
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=bold)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)


def _append_section_title(ws, row, title):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(name="Microsoft YaHei", size=12, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="305496")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    return row + 1


def _append_key_value_table(ws, row, rows):
    ws.cell(row=row, column=1, value="项目")
    ws.cell(row=row, column=2, value="内容")
    _style_range(ws, f"A{row}:B{row}", fill="D9EAF7", bold=True, align="center")
    row += 1
    start = row
    for key, value in rows:
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=value)
        row += 1
    _style_range(ws, f"A{start}:B{row - 1}")
    return row + 1


def _append_order_check_table(ws, row, demand, produced_totals):
    headers = ["产品编号", "规格 mm", "订单数量", "本单产出", "差异", "处理说明"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=header)
    _style_range(ws, f"A{row}:F{row}", fill="D9EAF7", bold=True, align="center")
    row += 1
    start = row

    for idx, (width, length, demand_num) in enumerate(
        zip(demand["Width"], demand["Length"], demand["num"]),
        start=1,
    ):
        produced = int(produced_totals.get(idx - 1, 0))
        gap = produced - int(demand_num)
        if gap == 0:
            note = "与订单一致"
        elif gap > 0:
            note = f"多 {gap} 件，作为余量"
        else:
            note = f"需补切 {abs(gap)} 件"
        values = [
            f"T{idx}",
            f"{_format_mm(width)} x {_format_mm(length)}",
            int(demand_num),
            produced,
            gap,
            note,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    _style_range(ws, f"A{start}:F{row - 1}")
    _style_range(ws, f"C{start}:E{row - 1}", align="right")
    return row + 1


def _append_stage_output_table(ws, row, output_rows):
    headers = ["产品编号", "每道产出", "本阶段产出", "产出重量(吨)", "单道尾料(mm)", "说明"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=header)
    _style_range(ws, f"A{row}:F{row}", fill="E2F0D9", bold=True, align="center")
    row += 1
    start = row
    for values in output_rows:
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1
    _style_range(ws, f"A{start}:F{row - 1}")
    _style_range(ws, f"B{start}:E{row - 1}", align="right")
    return row + 1


def _append_note_table(ws, row, rows):
    headers = ["项目", "说明"]
    ws.cell(row=row, column=1, value=headers[0])
    ws.cell(row=row, column=2, value=headers[1])
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    _style_range(ws, f"A{row}:F{row}", fill="D9EAF7", bold=True, align="center")
    row += 1
    start = row
    for key, value in rows:
        ws.cell(row=row, column=1, value=key)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.cell(row=row, column=2, value=value)
        row += 1
    _style_range(ws, f"A{start}:F{row - 1}")
    return row + 1


def _append_stage_sections(ws, row, demand, decoder_mode, solution):
    produced_totals = defaultdict(int)
    weight_unit = _infer_weight_unit(demand)

    if decoder_mode == "stage_based":
        stage_iter = [(strip, None) for strip in solution["strips"]]
    else:
        merged_strips, repeat_counts = merge_same_pattern_strips(solution["strips"])
        stage_iter = list(zip(merged_strips, repeat_counts))

    for idx, (strip, repeat_count) in enumerate(stage_iter, start=1):
        groups = _collect_lane_groups(strip)
        used_width = sum(group["width"] * group["lane_count"] for group in groups)
        waste_width = PANEL_WIDTH - used_width

        if decoder_mode == "stage_based":
            actual_length = float(strip.strip_length) * SCALE_FACTOR
            length_label = f"{_format_m(actual_length)} m"
            title_products = " + ".join(sorted({f"T{group['type_id'] + 1}" for group in groups}))
            title = f"阶段 {idx}：{title_products} 同一组刀位连续走料"
            rows = [
                ("排刀组合", _format_width_combo(groups)),
                ("刀数", f"{sum(group['lane_count'] for group in groups)} 道"),
                ("占用宽度", f"{int(used_width)} mm"),
                ("边部余宽", f"{int(waste_width)} mm"),
                ("定尺长度", " / ".join(f"T{group['type_id'] + 1}: {_format_mm(group['length'])} mm" for group in groups)),
                ("走料长度", length_label),
                ("预计用料重量", f"{_format_ton(_estimate_consumed_weight_tons(PANEL_WIDTH, actual_length, demand, weight_unit))} 吨"),
                ("刀位累计位置", f"{_format_knife_positions(groups)} mm"),
            ]
        else:
            actual_repeat = int(repeat_count) * SCALE_FACTOR
            actual_length = float(strip.strip_length) * actual_repeat
            title = f"阶段 {idx}：连续走料模式"
            rows = [
                ("排刀组合", _format_width_combo(groups)),
                ("刀数", f"{sum(group['lane_count'] for group in groups)} 道"),
                ("占用宽度", f"{int(used_width)} mm"),
                ("边部余宽", f"{int(waste_width)} mm"),
                ("单条长度", f"{_format_mm(strip.strip_length)} mm"),
                ("连续条数", f"{actual_repeat} 条"),
                ("走料长度", f"{_format_m(actual_length)} m"),
                ("预计用料重量", f"{_format_ton(_estimate_consumed_weight_tons(PANEL_WIDTH, actual_length, demand, weight_unit))} 吨"),
                ("刀位累计位置", f"{_format_knife_positions(groups)} mm"),
            ]

        row = _append_section_title(ws, row, title)
        row = _append_key_value_table(ws, row, rows)

        output_rows = []
        for group in groups:
            if decoder_mode == "stage_based":
                pieces_per_lane = int(actual_length // group["length"])
                tail_length = actual_length - pieces_per_lane * group["length"]
            else:
                pieces_per_lane = int(repeat_count) * SCALE_FACTOR
                tail_length = max(0, float(strip.strip_length) - float(group["length"]))
            lane_total = int(pieces_per_lane * group["lane_count"])
            produced_totals[group["type_id"]] += lane_total
            product_weight = lane_total * _piece_weight_tons(demand, group["type_id"], weight_unit)
            output_rows.append(
                [
                    f"T{group['type_id'] + 1}",
                    pieces_per_lane,
                    lane_total,
                    round(product_weight, 3),
                    round(float(tail_length), 2),
                    f"{group['lane_count']} 道并切",
                ]
            )
        row = _append_stage_output_table(ws, row, output_rows)
        row = _append_note_table(
            ws,
            row,
            [("示意图", f"母卷宽度 {PANEL_WIDTH} mm | {_build_stage_sketch(groups)}")],
        )

    return row, dict(produced_totals)


def write_cutting_report_xlsx(sheet_num, decoder_mode, demand, solution, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    production_summary = summarize_production(demand, solution, decoder_mode)
    workbook = Workbook()
    ws = workbook.active
    ws.title = "生产指令单"
    ws.sheet_view.showGridLines = False

    for column, width in {
        "A": 14,
        "B": 34,
        "C": 14,
        "D": 14,
        "E": 14,
        "F": 34,
    }.items():
        ws.column_dimensions[column].width = width

    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = "钢板纵切生产指令单"
    title.font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    row = 3
    row = _append_section_title(ws, row, "1. 生产任务信息")
    thickness = _parse_thickness_from_demand(demand)
    stage_count = len(solution["strips"]) if decoder_mode == "stage_based" else int(solution["num_strips"])
    row = _append_key_value_table(
        ws,
        row,
        [
            ("任务编号", sheet_num),
            ("类型", "分条"),
            ("母卷规格", f"{_format_mm(thickness) if thickness else '-'} x {PANEL_WIDTH} mm"),
            ("输入重量单位", "kg，报告已换算为吨" if production_summary["weight_unit"] == "kg" else "吨"),
            ("订单重量合计", f"{_format_ton(production_summary['input_weight_total_tons'])} 吨"),
            ("总走料长度", f"{_format_m(production_summary['real_total_length'])} m"),
            ("预计用料重量", f"{_format_ton(production_summary['consumed_weight_tons'])} 吨"),
            ("面积利用率", f"{production_summary['actual_utilization']:.2f}%"),
            ("生产阶段数", stage_count),
            ("解码方式", decoder_mode),
        ],
    )

    row = _append_section_title(ws, row, "2. 订单与产出核对")
    row = _append_order_check_table(ws, row, demand, production_summary["produced_totals"])

    row = _append_section_title(ws, row, "3. 排刀执行说明")
    row, _ = _append_stage_sections(ws, row, demand, decoder_mode, solution)

    shortage_notes = []
    over_notes = []
    for idx, demand_num in enumerate(demand["num"], start=1):
        produced = int(production_summary["produced_totals"].get(idx - 1, 0))
        gap = produced - int(demand_num)
        if gap < 0:
            shortage_notes.append(f"T{idx} 少 {abs(gap)} 件，需要单独补切")
        elif gap > 0:
            over_notes.append(f"T{idx} 多 {gap} 件，可作为余量件")

    row = _append_section_title(ws, row, "4. 换刀与补切提示")
    row = _append_note_table(
        ws,
        row,
        [
            ("换刀顺序", "按阶段编号顺序执行；每阶段内同一组刀位连续走料"),
            ("补切要求", "；".join(shortage_notes) if shortage_notes else "无"),
            ("多产说明", "；".join(over_notes) if over_notes else "无"),
            ("注意事项", "阶段执行过程中不要随意改变刀位组合"),
        ],
    )

    row = _append_section_title(ws, row, "5. 车间确认栏")
    row = _append_note_table(
        ws,
        row,
        [("排刀确认", ""), ("生产确认", ""), ("数量核对", ""), ("补切确认", ""), ("日期", "")],
    )

    ws.freeze_panes = "A3"
    workbook.save(output_path)
    return output_path


def build_cutting_report_text(sheet_num, decoder_mode, demand, solution):
    production_summary = summarize_production(demand, solution, decoder_mode)
    lines = []
    lines.extend(_format_task_section(sheet_num, decoder_mode, demand, solution, production_summary))
    lines.extend(_format_order_check_section(demand, production_summary["produced_totals"]))

    if decoder_mode == "stage_based":
        detail_lines, produced_totals = _build_stage_mode_section(demand, solution["strips"])
    else:
        merged_strips, repeat_counts = merge_same_pattern_strips(solution["strips"])
        detail_lines, produced_totals = _build_compact_mode_section(demand, merged_strips, repeat_counts)

    lines.extend(detail_lines)
    lines.extend(_format_execution_notes(demand, production_summary["produced_totals"]))
    lines.extend(_format_confirm_section())
    return "\n".join(lines)


def write_cutting_report(sheet_num, decoder_mode, demand, solution, output_path):
    output_path = Path(output_path)
    if output_path.suffix.lower() == ".xlsx":
        return write_cutting_report_xlsx(
            sheet_num=sheet_num,
            decoder_mode=decoder_mode,
            demand=demand,
            solution=solution,
            output_path=output_path,
        )

    output_path.write_text(
        build_cutting_report_text(sheet_num, decoder_mode, demand, solution),
        encoding="utf-8",
    )
    return output_path
