# -*- coding: utf-8 -*-
"""
Business workflow for the client's actual production workbook.

Current scope:
1. Read the adapted + routed group data.
2. Use direct plans for single-spec groups.
3. Convert multi-spec groups into the existing GA demand structure.
4. Run the current GA solver for multi-spec groups.
"""

from __future__ import annotations

from dataclasses import dataclass
from itertools import product
from pathlib import Path

import pandas as pd

from core.config import (
    AUTO_LARGE_BATCH_DECODER,
    AUTO_LARGE_BATCH_SCALE_FACTOR,
    AUTO_MEDIUM_BATCH_DECODER,
    AUTO_MEDIUM_BATCH_MAX_PIECES,
    AUTO_MEDIUM_BATCH_SCALE_FACTOR,
    AUTO_MULTI_SPEC_DECODER_MODE,
    AUTO_SMALL_BATCH_DECODER,
    AUTO_SMALL_BATCH_MAX_PIECES,
    AUTO_SMALL_BATCH_SCALE_FACTOR,
    CROSSOVER_RATE,
    ELITE_SIZE,
    MAX_GENERATIONS,
    MUTATION_RATE,
    PANEL_WIDTH,
    POPULATION_SIZE,
)
from core.group_routing import GroupRoutingResult
from core.production_state_adapter import AdaptedProductionData


MULTI_SPEC_ROTATION_SPEC_LIMIT = 2

WORKFLOW_SUMMARY_COLUMNS = [
    "分组编号",
    "品名",
    "厚度",
    "规格种数",
    "处理方式",
    "状态",
    "策略分类",
    "解码方式",
    "缩放因子",
    "母板宽度",
    "最优朝向",
    "订单总片数",
    "算法产出总片数",
    "补切数",
    "超产数",
    "总重量",
    "总消耗长度(mm)",
    "真实利用率(%)",
    "条带数/阶段数",
    "业务结论",
    "推荐动作",
    "图像输出",
    "报告输出",
    "备注",
]

UNSERVICEABLE_GROUP_COLUMNS = [
    "分组编号",
    "品名",
    "厚度",
    "规格种数",
    "当前母板宽度(mm)",
    "超宽规格",
    "最大短边(mm)",
    "建议处理",
]

MULTI_SPEC_RESULT_COLUMNS = [
    "分组编号",
    "品名",
    "厚度",
    "规格种数",
    "策略分类",
    "订单总片数",
    "算法产出总片数",
    "补切数",
    "超产数",
    "总重量",
    "解码方式",
    "缩放因子",
    "母板宽度",
    "编码后个体数",
    "编码后总长度(mm)",
    "还原后总长度(mm)",
    "真实利用率(%)",
    "条带数/阶段数",
    "惩罚值",
    "图像输出",
    "报告输出",
    "备注",
]

SINGLE_SPEC_RESULT_COLUMNS = [
    "分组编号",
    "品名",
    "厚度",
    "规格种数",
    "订单总片数",
    "算法产出总片数",
    "补切数",
    "超产数",
    "总重量",
    "母板宽度",
    "最优朝向",
    "最优横向占宽",
    "最优纵向定尺",
    "每条并排数",
    "需要条数",
    "总消耗长度(mm)",
    "余宽",
    "真实利用率(%)",
    "备注",
]


WF_GROUP_ID_COL = WORKFLOW_SUMMARY_COLUMNS[0]
WF_NAME_COL = WORKFLOW_SUMMARY_COLUMNS[1]
WF_THICKNESS_COL = WORKFLOW_SUMMARY_COLUMNS[2]
WF_SPEC_COUNT_COL = WORKFLOW_SUMMARY_COLUMNS[3]
WF_ROUTE_TYPE_COL = WORKFLOW_SUMMARY_COLUMNS[4]
WF_STATUS_COL = WORKFLOW_SUMMARY_COLUMNS[5]
WF_STRATEGY_BUCKET_COL = WORKFLOW_SUMMARY_COLUMNS[6]
WF_DECODER_MODE_COL = WORKFLOW_SUMMARY_COLUMNS[7]
WF_SCALE_FACTOR_COL = WORKFLOW_SUMMARY_COLUMNS[8]
WF_PANEL_WIDTH_COL = WORKFLOW_SUMMARY_COLUMNS[9]
WF_LAYOUT_NOTE_COL = WORKFLOW_SUMMARY_COLUMNS[10]
WF_ORDER_TOTAL_COL = WORKFLOW_SUMMARY_COLUMNS[11]
WF_PRODUCED_TOTAL_COL = WORKFLOW_SUMMARY_COLUMNS[12]
WF_MAKEUP_TOTAL_COL = WORKFLOW_SUMMARY_COLUMNS[13]
WF_OVER_TOTAL_COL = WORKFLOW_SUMMARY_COLUMNS[14]
WF_WEIGHT_TOTAL_COL = WORKFLOW_SUMMARY_COLUMNS[15]
WF_CONSUMED_LENGTH_COL = WORKFLOW_SUMMARY_COLUMNS[16]
WF_UTILIZATION_COL = WORKFLOW_SUMMARY_COLUMNS[17]
WF_STRIP_COUNT_COL = WORKFLOW_SUMMARY_COLUMNS[18]
WF_BUSINESS_DECISION_COL = WORKFLOW_SUMMARY_COLUMNS[19]
WF_RECOMMENDED_ACTION_COL = WORKFLOW_SUMMARY_COLUMNS[20]
WF_IMAGE_PATH_COL = WORKFLOW_SUMMARY_COLUMNS[21]
WF_REPORT_PATH_COL = WORKFLOW_SUMMARY_COLUMNS[22]
WF_NOTE_COL = WORKFLOW_SUMMARY_COLUMNS[23]

MS_GROUP_ID_COL = MULTI_SPEC_RESULT_COLUMNS[0]
MS_NAME_COL = MULTI_SPEC_RESULT_COLUMNS[1]
MS_THICKNESS_COL = MULTI_SPEC_RESULT_COLUMNS[2]
MS_SPEC_COUNT_COL = MULTI_SPEC_RESULT_COLUMNS[3]
MS_STRATEGY_BUCKET_COL = MULTI_SPEC_RESULT_COLUMNS[4]
MS_ORDER_TOTAL_COL = MULTI_SPEC_RESULT_COLUMNS[5]
MS_PRODUCED_TOTAL_COL = MULTI_SPEC_RESULT_COLUMNS[6]
MS_MAKEUP_TOTAL_COL = MULTI_SPEC_RESULT_COLUMNS[7]
MS_OVER_TOTAL_COL = MULTI_SPEC_RESULT_COLUMNS[8]
MS_WEIGHT_TOTAL_COL = MULTI_SPEC_RESULT_COLUMNS[9]
MS_DECODER_MODE_COL = MULTI_SPEC_RESULT_COLUMNS[10]
MS_SCALE_FACTOR_COL = MULTI_SPEC_RESULT_COLUMNS[11]
MS_PANEL_WIDTH_COL = MULTI_SPEC_RESULT_COLUMNS[12]
MS_EXPANDED_ITEM_COUNT_COL = MULTI_SPEC_RESULT_COLUMNS[13]
MS_RAW_CONSUMED_LENGTH_COL = MULTI_SPEC_RESULT_COLUMNS[14]
MS_REAL_CONSUMED_LENGTH_COL = MULTI_SPEC_RESULT_COLUMNS[15]
MS_UTILIZATION_COL = MULTI_SPEC_RESULT_COLUMNS[16]
MS_STRIP_COUNT_COL = MULTI_SPEC_RESULT_COLUMNS[17]
MS_PENALTY_COL = MULTI_SPEC_RESULT_COLUMNS[18]
MS_IMAGE_PATH_COL = MULTI_SPEC_RESULT_COLUMNS[19]
MS_REPORT_PATH_COL = MULTI_SPEC_RESULT_COLUMNS[20]
MS_NOTE_COL = MULTI_SPEC_RESULT_COLUMNS[21]

SS_GROUP_ID_COL = SINGLE_SPEC_RESULT_COLUMNS[0]
SS_PANEL_WIDTH_COL = SINGLE_SPEC_RESULT_COLUMNS[9]
SS_CONSUMED_LENGTH_COL = SINGLE_SPEC_RESULT_COLUMNS[15]
SS_UTILIZATION_COL = SINGLE_SPEC_RESULT_COLUMNS[17]

WIDTH_ALLOCATION_SUMMARY_COLUMNS = [
    "母板宽度",
    "分组数",
    "单规格组数",
    "多规格组数",
    "订单数",
    "算法产出数",
    "补切数",
    "超产数",
    "母板消耗长度(mm)",
    "母板面积",
    "产出面积",
    "真实利用率(%)",
]


@dataclass
class WorkflowExecutionResult:
    workflow_summary: pd.DataFrame
    single_spec_results: pd.DataFrame
    multi_spec_results: pd.DataFrame
    unserviceable_groups: pd.DataFrame
    width_allocation_summary: pd.DataFrame


@dataclass
class WorkflowRunSummary:
    panel_width: int
    decoder_mode: str
    total_groups: int
    completed_groups: int
    failed_groups: int
    single_spec_groups: int
    multi_spec_groups: int
    total_demand_area: float
    total_output_area: float
    total_makeup_pieces: int
    total_overproduction_pieces: int
    total_consumed_length_mm: float
    total_panel_area: float
    overall_efficiency: float


def load_runtime_modules():
    from core.cutting_report import write_cutting_report
    from core.data_loader import expand_demand
    import core.data_loader as data_loader_module
    import core.decoder as decoder_module
    from core.decoder import merge_same_pattern_strips
    from core.ga_engine_fast import GeneticAlgorithm
    import core.visualization as visualization_module
    from core.visualization import (
        plot_compact_cutting_plan,
        plot_stage_based_cutting_plan,
    )
    import core.cutting_report as cutting_report_module

    return {
        "write_cutting_report": write_cutting_report,
        "expand_demand": expand_demand,
        "data_loader_module": data_loader_module,
        "decoder_module": decoder_module,
        "visualization_module": visualization_module,
        "cutting_report_module": cutting_report_module,
        "merge_same_pattern_strips": merge_same_pattern_strips,
        "GeneticAlgorithm": GeneticAlgorithm,
        "plot_compact_cutting_plan": plot_compact_cutting_plan,
        "plot_stage_based_cutting_plan": plot_stage_based_cutting_plan,
    }


def apply_runtime_panel_width(runtime: dict, panel_width: int) -> None:
    import core.config as config_module

    config_module.PANEL_WIDTH = int(panel_width)
    runtime["decoder_module"].PANEL_WIDTH = int(panel_width)
    runtime["visualization_module"].PANEL_WIDTH = int(panel_width)
    runtime["cutting_report_module"].PANEL_WIDTH = int(panel_width)


def apply_runtime_scale_factor(runtime: dict, scale_factor: int) -> None:
    import core.config as config_module

    config_module.SCALE_FACTOR = int(scale_factor)
    runtime["data_loader_module"].SCALE_FACTOR = int(scale_factor)
    runtime["visualization_module"].SCALE_FACTOR = int(scale_factor)
    runtime["cutting_report_module"].SCALE_FACTOR = int(scale_factor)


def choose_multi_spec_strategy(group_row: dict, requested_decoder_mode: str) -> dict:
    total_pieces = int(group_row["总片数"])

    if total_pieces <= AUTO_SMALL_BATCH_MAX_PIECES:
        strategy_bucket = "小批量"
        auto_decoder_mode = AUTO_SMALL_BATCH_DECODER
        scale_factor = AUTO_SMALL_BATCH_SCALE_FACTOR
    elif total_pieces <= AUTO_MEDIUM_BATCH_MAX_PIECES:
        strategy_bucket = "中小批量"
        auto_decoder_mode = AUTO_MEDIUM_BATCH_DECODER
        scale_factor = AUTO_MEDIUM_BATCH_SCALE_FACTOR
    else:
        strategy_bucket = "大批量"
        auto_decoder_mode = AUTO_LARGE_BATCH_DECODER
        scale_factor = AUTO_LARGE_BATCH_SCALE_FACTOR

    manual_override = (
        requested_decoder_mode
        and str(requested_decoder_mode).lower() != AUTO_MULTI_SPEC_DECODER_MODE
    )
    selected_decoder_mode = requested_decoder_mode if manual_override else auto_decoder_mode

    strategy_note = (
        f"{strategy_bucket}自动策略：{selected_decoder_mode} + 缩放{scale_factor}"
        if not manual_override
        else f"{strategy_bucket}手动解码：{selected_decoder_mode} + 缩放{scale_factor}"
    )

    return {
        "strategy_bucket": strategy_bucket,
        "decoder_mode": str(selected_decoder_mode).lower(),
        "scale_factor": int(scale_factor),
        "strategy_note": strategy_note,
    }


def _derive_business_decision(row: pd.Series) -> pd.Series:
    status = str(row.get("状态", "") or "")
    note = str(row.get("备注", "") or "")
    makeup = 0 if pd.isna(row.get("补切数")) else int(row.get("补切数"))
    over = 0 if pd.isna(row.get("超产数")) else int(row.get("超产数"))

    if status != "已完成":
        if ("无法放入横向占宽" in note) or ("无法承接" in note) or ("外协" in note) or ("改规格" in note):
            return pd.Series(
                {
                    "业务结论": "需人工/外协/改规格处理",
                    "推荐动作": "当前母板无法承接，转人工评估、外协或改规格处理",
                }
            )
        return pd.Series(
            {
                "业务结论": "待人工复核",
                "推荐动作": "检查失败原因并人工复核后重新排版",
            }
        )

    if makeup > 0:
        return pd.Series(
            {
                "业务结论": "需补切后下发",
                "推荐动作": f"按当前方案先生产主体数量，另补切 {makeup} 件",
            }
        )

    if over > 0:
        return pd.Series(
            {
                "业务结论": "可直接下发",
                "推荐动作": f"按当前方案生产，并确认超产 {over} 件的库存去向",
            }
        )

    return pd.Series(
        {
            "业务结论": "可直接下发",
            "推荐动作": "按当前方案直接下发生产",
        }
    )


def attach_business_decisions(workflow_summary: pd.DataFrame) -> pd.DataFrame:
    if workflow_summary.empty:
        result = workflow_summary.copy()
        result["业务结论"] = pd.Series(dtype="object")
        result["推荐动作"] = pd.Series(dtype="object")
        return result

    result = workflow_summary.copy()
    decision_df = result.apply(_derive_business_decision, axis=1)
    result["业务结论"] = decision_df["业务结论"]
    result["推荐动作"] = decision_df["推荐动作"]
    return result


def _build_orientation_note(orientation_choices: dict[str, dict] | None) -> str:
    if not orientation_choices:
        return "短边横放"
    return "；".join(
        f"{spec}:{choice['orientation']}"
        for spec, choice in sorted(orientation_choices.items())
    )


def _build_multi_spec_orientation_candidates(
    group_detail: pd.DataFrame,
    panel_width: int,
) -> list[dict]:
    detail = group_detail.sort_values(["短边", "长边", "标准规格"], ignore_index=True)
    spec_count = int(detail["标准规格"].nunique())

    if spec_count != MULTI_SPEC_ROTATION_SPEC_LIMIT:
        return [
            {
                "orientation_choices": None,
                "orientation_note": "短边横放",
                "rotation_enabled": False,
            }
        ]

    per_spec_candidates: list[list[dict]] = []
    for row in detail.to_dict(orient="records"):
        spec = str(row["标准规格"])
        short_side = float(row["短边"])
        long_side = float(row["长边"])
        candidates = []

        if short_side <= float(panel_width):
            candidates.append(
                {
                    "spec": spec,
                    "width": short_side,
                    "length": long_side,
                    "orientation": "短边横放",
                }
            )

        if long_side <= float(panel_width) and (long_side, short_side) != (short_side, long_side):
            candidates.append(
                {
                    "spec": spec,
                    "width": long_side,
                    "length": short_side,
                    "orientation": "长边横放",
                }
            )

        if not candidates:
            raise ValueError(
                f"当前母板宽度 {panel_width} mm 下规格 {spec} 两个朝向均无法放入"
            )
        per_spec_candidates.append(candidates)

    orientation_candidates = []
    seen_keys: set[tuple] = set()
    for combo in product(*per_spec_candidates):
        key = tuple(
            (choice["spec"], float(choice["width"]), float(choice["length"]))
            for choice in combo
        )
        if key in seen_keys:
            continue
        seen_keys.add(key)

        orientation_choices = {choice["spec"]: choice for choice in combo}
        orientation_candidates.append(
            {
                "orientation_choices": orientation_choices,
                "orientation_note": _build_orientation_note(orientation_choices),
                "rotation_enabled": True,
            }
        )

    return orientation_candidates


def build_demand_from_group_detail(
    group_detail: pd.DataFrame,
    orientation_choices: dict[str, dict] | None = None,
) -> dict:
    detail = group_detail.sort_values(["短边", "长边", "标准规格"], ignore_index=True)
    widths = []
    lengths = []
    orientations = []
    specs = []

    for row in detail.to_dict(orient="records"):
        spec = str(row["标准规格"])
        choice = (orientation_choices or {}).get(spec)
        if choice is None:
            width = float(row["短边"])
            length = float(row["长边"])
            orientation = "短边横放"
        else:
            width = float(choice["width"])
            length = float(choice["length"])
            orientation = str(choice["orientation"])
        widths.append(width)
        lengths.append(length)
        orientations.append(orientation)
        specs.append(spec)

    return {
        "Width": widths,
        "Length": lengths,
        "num": [int(value) for value in detail["片数"]],
        "Weight": [float(value) for value in detail["重量"]],
        "Orientation": orientations,
        "Spec": specs,
        "Name": str(detail.iloc[0, 1]) if not detail.empty else "",
    }


def _summarize_group_production(runtime: dict, demand: dict, solution: dict, decoder_mode: str) -> dict:
    return runtime["cutting_report_module"].summarize_production(
        demand=demand,
        solution=solution,
        decoder_mode=decoder_mode,
    )


def _build_unserviceable_groups(
    adapted_data: AdaptedProductionData,
    workflow_summary: pd.DataFrame,
) -> pd.DataFrame:
    failed_rows = workflow_summary.loc[
        workflow_summary["备注"].astype(str).str.contains("无法放入横向占宽", na=False),
        ["分组编号", "品名", "厚度", "规格种数", "母板宽度"],
    ].copy()

    rows: list[dict] = []
    for group_row in failed_rows.to_dict(orient="records"):
        panel_width = float(group_row["母板宽度"])
        group_detail = adapted_data.printable_detail.loc[
            adapted_data.printable_detail["分组编号"] == group_row["分组编号"]
        ].copy()
        oversize_detail = group_detail.loc[group_detail["短边"] > panel_width].copy()

        oversize_specs = sorted(oversize_detail["标准规格"].astype(str).unique().tolist())
        max_short_side = float(oversize_detail["短边"].max()) if not oversize_detail.empty else None

        rows.append(
            {
                "分组编号": group_row["分组编号"],
                "品名": group_row["品名"],
                "厚度": group_row["厚度"],
                "规格种数": group_row["规格种数"],
                "当前母板宽度(mm)": panel_width,
                "超宽规格": "；".join(oversize_specs),
                "最大短边(mm)": max_short_side,
                "建议处理": "现有母板无法承接，需人工/外协/改规格处理",
            }
        )

    unserviceable_groups = pd.DataFrame(rows)
    if unserviceable_groups.empty:
        return pd.DataFrame(columns=UNSERVICEABLE_GROUP_COLUMNS)

    return unserviceable_groups.loc[:, UNSERVICEABLE_GROUP_COLUMNS].sort_values(
        ["最大短边(mm)", "分组编号"],
        ascending=[False, True],
        ignore_index=True,
    )


def _save_multi_spec_outputs(
    runtime: dict,
    group_id: str,
    decoder_mode: str,
    demand: dict,
    result: dict,
    output_dir: Path,
    panel_width: int,
) -> tuple[Path, Path]:
    solution = result["solution"]
    efficiency = result["real_efficiency"]

    # Keep visualization/report generation aligned with the selected panel width.
    apply_runtime_panel_width(runtime, panel_width)

    image_path = output_dir / f"{group_id}_{decoder_mode}_cutting_plan.png"
    report_path = output_dir / f"{group_id}_{decoder_mode}_cutting_report.txt"

    if decoder_mode == "stage_based":
        runtime["plot_stage_based_cutting_plan"](
            solution["strips"],
            efficiency,
            save_path=image_path,
            show=False,
        )
    else:
        merged_strips, repeat_counts = runtime["merge_same_pattern_strips"](solution["strips"])
        runtime["plot_compact_cutting_plan"](
            merged_strips,
            repeat_counts,
            efficiency,
            save_path=image_path,
            show=False,
        )

    runtime["write_cutting_report"](
        sheet_num=group_id,
        decoder_mode=decoder_mode,
        demand=demand,
        solution=solution,
        output_path=report_path,
    )

    return image_path, report_path


def run_multi_spec_group(
    group_row: dict,
    group_detail: pd.DataFrame,
    decoder_mode: str,
    runtime: dict,
    output_dir: Path,
    panel_width: int,
    save_outputs: bool = True,
    orientation_choices: dict[str, dict] | None = None,
    orientation_note: str | None = None,
) -> dict:
    if group_detail.empty:
        raise ValueError("分组缺少可排版明细")

    demand = build_demand_from_group_detail(
        group_detail,
        orientation_choices=orientation_choices,
    )
    oversize_widths = sorted({int(width) for width in demand["Width"] if float(width) > float(panel_width)})
    if oversize_widths:
        width_text = ", ".join(str(width) for width in oversize_widths)
        raise ValueError(
            f"当前母板宽度 {panel_width} mm 下无法放入横向占宽 {width_text} mm 的规格"
        )

    strategy = choose_multi_spec_strategy(
        group_row=group_row,
        requested_decoder_mode=decoder_mode,
    )
    apply_runtime_scale_factor(runtime, strategy["scale_factor"])

    items, _ = runtime["expand_demand"](demand)
    if not items:
        raise ValueError("展开后无可用排版个体")

    ga = runtime["GeneticAlgorithm"](
        items,
        population_size=POPULATION_SIZE,
        max_generations=MAX_GENERATIONS,
        crossover_rate=CROSSOVER_RATE,
        mutation_rate=MUTATION_RATE,
        elite_size=ELITE_SIZE,
        decoder_mode=strategy["decoder_mode"],
    )
    best_individual, best_fitness = ga.evolve(verbose=False)
    solution = ga.get_solution(best_individual)
    if int(solution["num_strips"]) <= 0 or float(solution["total_length"]) <= 0:
        raise ValueError("当前母板宽度下未生成有效排版方案")

    production_summary = _summarize_group_production(
        runtime=runtime,
        demand=demand,
        solution=solution,
        decoder_mode=strategy["decoder_mode"],
    )

    if save_outputs:
        image_path, report_path = _save_multi_spec_outputs(
            runtime=runtime,
            group_id=group_row["分组编号"],
            decoder_mode=strategy["decoder_mode"],
            demand=demand,
            result={
                "solution": solution,
                "real_efficiency": production_summary["actual_utilization"],
            },
            output_dir=output_dir,
            panel_width=panel_width,
        )
    else:
        image_path = Path("")
        report_path = Path("")

    if int(group_row["规格种数"]) <= 1:
        result_note = (
            f"单规格无直排方案，按 {strategy['decoder_mode']} 模式转入 GA 兜底；"
            f"{strategy['strategy_note']}"
        )
    else:
        result_note = f"按 {strategy['decoder_mode']} 模式完成 GA 求解；{strategy['strategy_note']}"
    if orientation_note:
        result_note = f"{result_note}；排版朝向：{orientation_note}"

    return {
        "分组编号": group_row["分组编号"],
        "品名": group_row["品名"],
        "厚度": group_row["厚度"],
        "规格种数": group_row["规格种数"],
        "策略分类": strategy["strategy_bucket"],
        "订单总片数": int(production_summary["order_total"]),
        "算法产出总片数": int(production_summary["produced_total"]),
        "补切数": int(production_summary["shortage_total"]),
        "超产数": int(production_summary["over_total"]),
        "总重量": group_row["总重量"],
        "解码方式": strategy["decoder_mode"],
        "缩放因子": int(strategy["scale_factor"]),
        "母板宽度": panel_width,
        "编码后个体数": len(items),
        "编码后总长度(mm)": float(solution["total_length"]),
        "还原后总长度(mm)": float(production_summary["real_total_length"]),
        "真实利用率(%)": float(production_summary["actual_utilization"]),
        "条带数/阶段数": int(solution["num_strips"]),
        "惩罚值": float(solution["penalty"]),
        "图像输出": str(image_path),
        "报告输出": str(report_path),
        "备注": result_note,
        "__orientation_note__": orientation_note or _build_orientation_note(orientation_choices),
        "__save_context__": {
            "group_id": group_row["分组编号"],
            "decoder_mode": strategy["decoder_mode"],
            "demand": demand,
            "solution": solution,
            "real_efficiency": production_summary["actual_utilization"],
        },
    }


def _score_multi_spec_result(result_row: dict) -> tuple:
    panel_width = float(result_row[MS_PANEL_WIDTH_COL])
    consumed_length = float(result_row[MS_REAL_CONSUMED_LENGTH_COL])
    panel_area = panel_width * consumed_length
    utilization = float(result_row[MS_UTILIZATION_COL])
    return (
        int(result_row[MS_MAKEUP_TOTAL_COL]),
        int(result_row[MS_OVER_TOTAL_COL]),
        panel_area,
        consumed_length,
        -round(utilization, 8),
        panel_width,
    )


def _build_width_allocation_summary(workflow_summary: pd.DataFrame) -> pd.DataFrame:
    if workflow_summary.empty:
        return pd.DataFrame(columns=WIDTH_ALLOCATION_SUMMARY_COLUMNS)

    completed_rows = workflow_summary.loc[
        workflow_summary[WF_PRODUCED_TOTAL_COL].notna() & workflow_summary[WF_PANEL_WIDTH_COL].notna()
    ].copy()
    if completed_rows.empty:
        return pd.DataFrame(columns=WIDTH_ALLOCATION_SUMMARY_COLUMNS)

    completed_rows["_is_multi_spec"] = completed_rows[WF_DECODER_MODE_COL].fillna("").astype(str).str.len() > 0
    completed_rows["_is_single_spec"] = ~completed_rows["_is_multi_spec"]
    completed_rows["_panel_area"] = (
        completed_rows[WF_PANEL_WIDTH_COL].astype(float)
        * completed_rows[WF_CONSUMED_LENGTH_COL].fillna(0).astype(float)
    )
    completed_rows["_output_area"] = (
        completed_rows["_panel_area"]
        * completed_rows[WF_UTILIZATION_COL].fillna(0).astype(float)
        / 100.0
    )

    summary_df = (
        completed_rows.groupby(WF_PANEL_WIDTH_COL, dropna=False)
        .agg(
            分组数=(WF_GROUP_ID_COL, "count"),
            单规格组数=("_is_single_spec", "sum"),
            多规格组数=("_is_multi_spec", "sum"),
            订单数=(WF_ORDER_TOTAL_COL, "sum"),
            算法产出数=(WF_PRODUCED_TOTAL_COL, "sum"),
            补切数=(WF_MAKEUP_TOTAL_COL, "sum"),
            超产数=(WF_OVER_TOTAL_COL, "sum"),
            **{
                "母板消耗长度(mm)": (WF_CONSUMED_LENGTH_COL, "sum"),
                "母板面积": ("_panel_area", "sum"),
                "产出面积": ("_output_area", "sum"),
            },
        )
        .reset_index()
        .rename(columns={WF_PANEL_WIDTH_COL: "母板宽度"})
    )

    summary_df["真实利用率(%)"] = 0.0
    valid_mask = summary_df["母板面积"] > 0
    summary_df.loc[valid_mask, "真实利用率(%)"] = (
        100.0 * summary_df.loc[valid_mask, "产出面积"] / summary_df.loc[valid_mask, "母板面积"]
    )

    integer_columns = [
        "母板宽度",
        "分组数",
        "单规格组数",
        "多规格组数",
        "订单数",
        "算法产出数",
        "补切数",
        "超产数",
    ]
    for column in integer_columns:
        summary_df[column] = summary_df[column].astype(int)

    return summary_df.loc[:, WIDTH_ALLOCATION_SUMMARY_COLUMNS].sort_values(
        ["母板宽度"],
        ignore_index=True,
    )


def _build_single_spec_report_text(
    *,
    group_id: str,
    name: str,
    thickness: float,
    panel_width: float,
    orientation: str,
    layout_width: float,
    layout_length: float,
    lane_count: int,
    strip_count: int,
    order_total: int,
    produced_total: int,
    makeup_total: int,
    over_total: int,
    weight_total: float,
    consumed_length: float,
    waste_width: float,
    utilization: float,
) -> str:
    input_weight_tons = float(weight_total) / 1000.0
    consumed_weight_tons = (
        float(panel_width)
        * float(consumed_length)
        * float(thickness)
        * 7.85e-9
    )
    product_weight_tons = (
        input_weight_tons * float(produced_total) / float(order_total)
        if int(order_total) > 0
        else 0.0
    )
    lines = [
        "1. 生产任务信息",
        "项目      内容",
        f"任务编号    {group_id}",
        f"品名      {name}",
        f"母卷规格    {thickness:g} x {int(panel_width)} mm",
        "输入重量单位  kg，报告已换算为吨",
        f"订单重量合计  {input_weight_tons:.6f} 吨",
        f"总走料长度  {float(consumed_length) / 1000.0:.3f} m",
        f"预计用料重量  {consumed_weight_tons:.6f} 吨",
        f"面积利用率  {float(utilization):.2f}%",
        "生产阶段数  1",
        "",
        "2. 订单与产出核对",
        "产品编号 规格 mm              订单数量   本单产出      差异 处理说明",
        (
            f"T1     {int(layout_width)} x {int(layout_length):<12} "
            f"{int(order_total):>8} {int(produced_total):>8} "
            f"{int(produced_total) - int(order_total):>7} "
            f"{'与订单一致' if int(produced_total) == int(order_total) else ('需补切 ' + str(int(makeup_total)) + ' 件' if int(makeup_total) > 0 else '多 ' + str(int(over_total)) + ' 件，作为余量')}"
        ),
        "",
        "3. 排刀执行说明",
        "阶段 1：T1 单规格直排",
        "项目      内容",
        f"排刀组合    {' + '.join(str(int(layout_width)) for _ in range(int(lane_count)))}",
        f"刀数      {int(lane_count)} 道",
        f"占用宽度    {int(layout_width * lane_count)} mm",
        f"边部余宽    {int(waste_width)} mm",
        f"定尺长度    {int(layout_length)} mm",
        f"走料长度    {float(consumed_length) / 1000.0:.3f} m",
        f"预计用料重量  {consumed_weight_tons:.6f} 吨",
        f"每道产出    {int(strip_count)} 件",
        f"本阶段产出 T1：{int(produced_total)} 件，约 {product_weight_tons:.6f} 吨",
        f"示意图：母卷宽度 {int(panel_width)} mm | "
        f"{' | '.join(['T1 ' + str(int(layout_width)) for _ in range(int(lane_count))])}"
        + (f" | 余宽 {int(waste_width)}" if waste_width > 0 else ""),
        "",
        "4. 换刀与补切提示",
        "项目   说明",
        "换刀顺序 按阶段 1 执行",
        f"补切要求 {'无' if int(makeup_total) == 0 else 'T1 少 ' + str(int(makeup_total)) + ' 件，需要单独补切'}",
        f"多产说明 {'无' if int(over_total) == 0 else 'T1 多 ' + str(int(over_total)) + ' 件，可作为余量件'}",
        "注意事项 阶段执行过程中不要随意改变刀位组合",
        "",
        "5. 车间确认栏",
        "项目   签字 / 确认",
        "排刀确认",
        "生产确认",
        "数量核对",
        "补切确认",
        "日期",
    ]
    return "\n".join(lines) + "\n"


def _write_single_spec_report_xlsx(
    *,
    output_path: Path,
    group_id: str,
    name: str,
    thickness: float,
    panel_width: float,
    orientation: str,
    layout_width: float,
    layout_length: float,
    lane_count: int,
    strip_count: int,
    order_total: int,
    produced_total: int,
    makeup_total: int,
    over_total: int,
    weight_total: float,
    consumed_length: float,
    waste_width: float,
    utilization: float,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    def style_range(ws, cell_range, *, fill=None, bold=False, align="left"):
        side = Side(style="thin", color="B7B7B7")
        border = Border(left=side, right=side, top=side, bottom=side)
        for row_cells in ws[cell_range]:
            for cell in row_cells:
                cell.border = border
                cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
                cell.font = Font(name="Microsoft YaHei", size=10, bold=bold)
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)

    def section(ws, row, title):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(name="Microsoft YaHei", size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 24
        return row + 1

    def key_value(ws, row, rows):
        ws.cell(row=row, column=1, value="项目")
        ws.cell(row=row, column=2, value="内容")
        style_range(ws, f"A{row}:B{row}", fill="D9EAF7", bold=True, align="center")
        row += 1
        start = row
        for key, value in rows:
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=value)
            row += 1
        style_range(ws, f"A{start}:B{row - 1}")
        return row + 1

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    input_weight_tons = float(weight_total) / 1000.0
    consumed_weight_tons = (
        float(panel_width)
        * float(consumed_length)
        * float(thickness)
        * 7.85e-9
    )
    product_weight_tons = (
        input_weight_tons * float(produced_total) / float(order_total)
        if int(order_total) > 0
        else 0.0
    )
    gap = int(produced_total) - int(order_total)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "生产指令单"
    ws.sheet_view.showGridLines = False
    for column, width in {"A": 14, "B": 34, "C": 14, "D": 14, "E": 14, "F": 34}.items():
        ws.column_dimensions[column].width = width

    ws.merge_cells("A1:F1")
    title = ws["A1"]
    title.value = "钢板纵切生产指令单"
    title.font = Font(name="Microsoft YaHei", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    row = 3
    row = section(ws, row, "1. 生产任务信息")
    row = key_value(
        ws,
        row,
        [
            ("任务编号", group_id),
            ("类型", "分条"),
            ("材料名称", name),
            ("母卷规格", f"{thickness:g} x {int(panel_width)} mm"),
            ("输入重量单位", "kg，报告已换算为吨"),
            ("订单重量合计", f"{input_weight_tons:.6f} 吨"),
            ("总走料长度", f"{float(consumed_length) / 1000.0:.3f} m"),
            ("预计用料重量", f"{consumed_weight_tons:.6f} 吨"),
            ("面积利用率", f"{float(utilization):.2f}%"),
            ("生产阶段数", 1),
            ("排版朝向", orientation),
        ],
    )

    row = section(ws, row, "2. 订单与产出核对")
    headers = ["产品编号", "规格 mm", "订单数量", "本单产出", "差异", "处理说明"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=header)
    style_range(ws, f"A{row}:F{row}", fill="D9EAF7", bold=True, align="center")
    row += 1
    note = (
        "与订单一致"
        if gap == 0
        else (f"需补切 {int(makeup_total)} 件" if gap < 0 else f"多 {int(over_total)} 件，作为余量")
    )
    values = ["T1", f"{int(layout_width)} x {int(layout_length)}", int(order_total), int(produced_total), gap, note]
    for col, value in enumerate(values, start=1):
        ws.cell(row=row, column=col, value=value)
    style_range(ws, f"A{row}:F{row}")
    style_range(ws, f"C{row}:E{row}", align="right")
    row += 2

    row = section(ws, row, "3. 排刀执行说明")
    row = section(ws, row, "阶段 1：T1 单规格直排")
    row = key_value(
        ws,
        row,
        [
            ("排刀组合", " + ".join(str(int(layout_width)) for _ in range(int(lane_count)))),
            ("刀数", f"{int(lane_count)} 道"),
            ("占用宽度", f"{int(layout_width * lane_count)} mm"),
            ("边部余宽", f"{int(waste_width)} mm"),
            ("定尺长度", f"{int(layout_length)} mm"),
            ("走料长度", f"{float(consumed_length) / 1000.0:.3f} m"),
            ("预计用料重量", f"{consumed_weight_tons:.6f} 吨"),
            ("每道产出", f"{int(strip_count)} 件"),
            ("本阶段产出", f"T1：{int(produced_total)} 件，约 {product_weight_tons:.6f} 吨"),
            (
                "示意图",
                f"母卷宽度 {int(panel_width)} mm | "
                f"{' | '.join(['T1 ' + str(int(layout_width)) for _ in range(int(lane_count))])}"
                + (f" | 余宽 {int(waste_width)}" if waste_width > 0 else ""),
            ),
        ],
    )

    row = section(ws, row, "4. 换刀与补切提示")
    row = key_value(
        ws,
        row,
        [
            ("换刀顺序", "按阶段 1 执行"),
            ("补切要求", "无" if int(makeup_total) == 0 else f"T1 少 {int(makeup_total)} 件，需要单独补切"),
            ("多产说明", "无" if int(over_total) == 0 else f"T1 多 {int(over_total)} 件，可作为余量件"),
            ("注意事项", "阶段执行过程中不要随意改变刀位组合"),
        ],
    )

    row = section(ws, row, "5. 车间确认栏")
    key_value(ws, row, [("排刀确认", ""), ("生产确认", ""), ("数量核对", ""), ("补切确认", ""), ("日期", "")])

    ws.freeze_panes = "A3"
    workbook.save(output_path)
    return output_path


def _save_single_spec_outputs(
    *,
    group_id: str,
    name: str,
    thickness: float,
    panel_width: float,
    orientation: str,
    layout_width: float,
    layout_length: float,
    lane_count: int,
    strip_count: int,
    order_total: int,
    produced_total: int,
    makeup_total: int,
    over_total: int,
    weight_total: float,
    consumed_length: float,
    waste_width: float,
    utilization: float,
    output_dir: Path,
) -> tuple[Path, Path]:
    import matplotlib.pyplot as plt
    import matplotlib.patches as patches

    plt.rcParams["font.sans-serif"] = ["SimHei"]
    plt.rcParams["axes.unicode_minus"] = False

    output_dir.mkdir(parents=True, exist_ok=True)
    image_path = output_dir / f"{group_id}_single_spec_cutting_plan.png"
    report_path = output_dir / f"{group_id}_single_spec_cutting_report.txt"

    total_length = max(float(consumed_length), 1.0)
    fig, ax = plt.subplots(figsize=(12, 4.5))
    scale_y = 10.0 / max(float(panel_width), 1.0)
    scale_x = 18.0 / total_length

    bg_rect = patches.Rectangle(
        (0, 0),
        total_length * scale_x,
        panel_width * scale_y,
        linewidth=1.5,
        edgecolor="black",
        facecolor="#efefef",
    )
    ax.add_patch(bg_rect)

    y_cursor = 0.0
    for lane_idx in range(int(lane_count)):
        lane_rect = patches.Rectangle(
            (0, y_cursor),
            total_length * scale_x,
            layout_width * scale_y,
            linewidth=1.0,
            edgecolor="black",
            facecolor="#8ecae6",
            alpha=0.9,
        )
        ax.add_patch(lane_rect)
        ax.text(
            total_length * scale_x / 2,
            y_cursor + layout_width * scale_y / 2,
            f"Lane {lane_idx + 1}\n{int(layout_width)} x {int(layout_length)}\n{int(strip_count)} pcs/lane",
            ha="center",
            va="center",
            fontsize=8,
        )
        y_cursor += layout_width * scale_y

    if waste_width > 0:
        waste_rect = patches.Rectangle(
            (0, y_cursor),
            total_length * scale_x,
            waste_width * scale_y,
            linewidth=1.0,
            edgecolor="black",
            facecolor="#f4a261",
            alpha=0.7,
        )
        ax.add_patch(waste_rect)
        ax.text(
            total_length * scale_x / 2,
            y_cursor + waste_width * scale_y / 2,
            f"余宽 {int(waste_width)} mm",
            ha="center",
            va="center",
            fontsize=8,
        )

    ax.set_title(
        f"{group_id} 单规格直排 | 母板{int(panel_width)} | 产出{int(produced_total)} | 利用率{float(utilization):.2f}%",
        fontsize=11,
    )
    ax.set_xlim(-0.5, total_length * scale_x + 0.5)
    ax.set_ylim(0, panel_width * scale_y + 0.8)
    ax.axis("off")
    plt.tight_layout()
    plt.savefig(image_path, dpi=160, bbox_inches="tight")
    plt.close(fig)

    report_path.write_text(
        _build_single_spec_report_text(
            group_id=group_id,
            name=name,
            thickness=thickness,
            panel_width=panel_width,
            orientation=orientation,
            layout_width=layout_width,
            layout_length=layout_length,
            lane_count=lane_count,
            strip_count=strip_count,
            order_total=order_total,
            produced_total=produced_total,
            makeup_total=makeup_total,
            over_total=over_total,
            weight_total=weight_total,
            consumed_length=consumed_length,
            waste_width=waste_width,
            utilization=utilization,
        ),
        encoding="utf-8",
    )

    return image_path, report_path


def execute_actual_production_workflow(
    adapted_data: AdaptedProductionData,
    routing_result: GroupRoutingResult,
    decoder_mode: str,
    output_dir: str | Path,
    panel_width: int = PANEL_WIDTH,
    verbose: bool = True,
) -> WorkflowExecutionResult:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    runtime = load_runtime_modules()
    apply_runtime_panel_width(runtime, panel_width)
    apply_runtime_scale_factor(runtime, AUTO_LARGE_BATCH_SCALE_FACTOR)

    detail = adapted_data.printable_detail.copy()
    routing_summary = routing_result.routing_summary.copy()

    single_spec_results = routing_summary.loc[
        routing_summary["处理方式"] == "单规格直排",
        [
            "分组编号",
            "品名",
            "厚度",
            "规格种数",
            "总片数",
            "总重量",
            "最优母板宽度",
            "最优朝向",
            "最优横向占宽",
            "最优纵向定尺",
            "每条并排数",
            "需要条数",
            "实际产出",
            "超产数",
            "总消耗长度",
            "余宽",
            "利用率(%)",
            "备注",
        ],
    ].copy()
    if single_spec_results.empty:
        single_spec_results = pd.DataFrame(columns=SINGLE_SPEC_RESULT_COLUMNS)
    else:
        single_spec_results["订单总片数"] = single_spec_results["总片数"].astype(int)
        single_spec_results["算法产出总片数"] = single_spec_results["实际产出"].astype(int)
        single_spec_results["补切数"] = (
            single_spec_results["订单总片数"] - single_spec_results["算法产出总片数"]
        ).clip(lower=0)
        single_spec_results["超产数"] = (
            single_spec_results["算法产出总片数"] - single_spec_results["订单总片数"]
        ).clip(lower=0)
        consumed_area = (
            single_spec_results["最优母板宽度"].fillna(0)
            * single_spec_results["总消耗长度"].fillna(0)
        )
        produced_area = (
            single_spec_results["算法产出总片数"].fillna(0)
            * single_spec_results["最优横向占宽"].fillna(0)
            * single_spec_results["最优纵向定尺"].fillna(0)
        )
        single_spec_results["真实利用率(%)"] = 0.0
        valid_mask = consumed_area > 0
        single_spec_results.loc[valid_mask, "真实利用率(%)"] = (
            100.0 * produced_area.loc[valid_mask] / consumed_area.loc[valid_mask]
        )
        single_spec_results = single_spec_results.rename(
            columns={
                "最优母板宽度": "母板宽度",
                "总消耗长度": "总消耗长度(mm)",
            }
        )
        single_spec_results = single_spec_results.loc[:, SINGLE_SPEC_RESULT_COLUMNS]

    multi_spec_rows: list[dict] = []
    workflow_rows: list[dict] = []

    for group_row in routing_summary.to_dict(orient="records"):
        if group_row["处理方式"] == "单规格直排":
            workflow_rows.append(
                {
                    "分组编号": group_row["分组编号"],
                    "品名": group_row["品名"],
                    "厚度": group_row["厚度"],
                    "规格种数": group_row["规格种数"],
                    "处理方式": group_row["处理方式"],
                    "状态": "已完成",
                    "策略分类": "单规格直排",
                    "解码方式": "",
                    "缩放因子": None,
                    "母板宽度": group_row["最优母板宽度"],
                    "最优朝向": group_row["最优朝向"],
                    "订单总片数": int(group_row["总片数"]),
                    "算法产出总片数": int(group_row["实际产出"]),
                    "补切数": max(int(group_row["总片数"]) - int(group_row["实际产出"]), 0),
                    "超产数": max(int(group_row["实际产出"]) - int(group_row["总片数"]), 0),
                    "总重量": group_row["总重量"],
                    "总消耗长度(mm)": group_row["总消耗长度"],
                    "真实利用率(%)": (
                        100.0
                        * float(group_row["实际产出"])
                        * float(group_row["最优横向占宽"])
                        * float(group_row["最优纵向定尺"])
                        / (float(group_row["最优母板宽度"]) * float(group_row["总消耗长度"]))
                        if group_row["最优母板宽度"] and group_row["总消耗长度"]
                        else 0.0
                    ),
                    "条带数/阶段数": group_row["需要条数"],
                    "图像输出": "",
                    "报告输出": "",
                    "备注": group_row["备注"],
                }
            )
            continue

        group_detail = detail.loc[detail["分组编号"] == group_row["分组编号"]].copy()
        is_single_spec_fallback = int(group_row["规格种数"]) <= 1
        group_strategy = choose_multi_spec_strategy(
            group_row=group_row,
            requested_decoder_mode=decoder_mode,
        )
        if verbose:
            process_label = "单规格转GA" if is_single_spec_fallback else "多规格优化"
            print(
                f"[{process_label}] {group_row['分组编号']} "
                f"{group_row['品名']} 厚度 {group_row['厚度']} "
                f"规格种数 {group_row['规格种数']} "
                f"-> {group_strategy['strategy_note']}"
            )

        try:
            multi_result = run_multi_spec_group(
                group_row=group_row,
                group_detail=group_detail,
                decoder_mode=decoder_mode,
                runtime=runtime,
                output_dir=output_dir,
                panel_width=panel_width,
            )
        except Exception as exc:
            workflow_rows.append(
                {
                    "分组编号": group_row["分组编号"],
                    "品名": group_row["品名"],
                    "厚度": group_row["厚度"],
                    "规格种数": group_row["规格种数"],
                    "处理方式": "多规格优化",
                    "状态": "失败",
                    "策略分类": group_strategy["strategy_bucket"],
                    "解码方式": group_strategy["decoder_mode"],
                    "缩放因子": group_strategy["scale_factor"],
                    "母板宽度": panel_width,
                    "最优朝向": "",
                    "订单总片数": group_row["总片数"],
                    "算法产出总片数": None,
                    "补切数": None,
                    "超产数": None,
                    "总重量": group_row["总重量"],
                    "总消耗长度(mm)": None,
                    "真实利用率(%)": None,
                    "条带数/阶段数": None,
                    "图像输出": "",
                    "报告输出": "",
                    "备注": (
                        f"单规格直排失败后转 GA 仍失败: {exc}"
                        if is_single_spec_fallback
                        else str(exc)
                    ),
                }
            )
            if verbose:
                print(f"  -> 失败: {exc}")
            continue

        multi_spec_rows.append(multi_result)
        workflow_rows.append(
            {
                "分组编号": multi_result["分组编号"],
                "品名": multi_result["品名"],
                "厚度": multi_result["厚度"],
                "规格种数": multi_result["规格种数"],
                "处理方式": "多规格优化",
                "状态": "已完成",
                "策略分类": multi_result["策略分类"],
                "解码方式": multi_result["解码方式"],
                "缩放因子": multi_result["缩放因子"],
                "母板宽度": multi_result["母板宽度"],
                "最优朝向": "",
                "订单总片数": multi_result["订单总片数"],
                "算法产出总片数": multi_result["算法产出总片数"],
                "补切数": multi_result["补切数"],
                "超产数": multi_result["超产数"],
                "总重量": multi_result["总重量"],
                "总消耗长度(mm)": multi_result["还原后总长度(mm)"],
                "真实利用率(%)": multi_result["真实利用率(%)"],
                "条带数/阶段数": multi_result["条带数/阶段数"],
                "图像输出": multi_result["图像输出"],
                "报告输出": multi_result["报告输出"],
                "备注": multi_result["备注"],
            }
        )

    workflow_summary = pd.DataFrame(workflow_rows)
    if workflow_summary.empty:
        workflow_summary = pd.DataFrame(columns=WORKFLOW_SUMMARY_COLUMNS)
    else:
        workflow_summary = attach_business_decisions(workflow_summary)
        workflow_summary = workflow_summary.loc[:, WORKFLOW_SUMMARY_COLUMNS].sort_values(
            [WORKFLOW_SUMMARY_COLUMNS[0]],
            ignore_index=True,
        )

    multi_spec_results = pd.DataFrame(multi_spec_rows)
    if multi_spec_results.empty:
        multi_spec_results = pd.DataFrame(columns=MULTI_SPEC_RESULT_COLUMNS)
    else:
        multi_spec_results = multi_spec_results.loc[:, MULTI_SPEC_RESULT_COLUMNS].sort_values(
            [MULTI_SPEC_RESULT_COLUMNS[0]],
            ignore_index=True,
        )

    unserviceable_groups = _build_unserviceable_groups(
        adapted_data=adapted_data,
        workflow_summary=workflow_summary,
    )

    return WorkflowExecutionResult(
        workflow_summary=workflow_summary,
        single_spec_results=single_spec_results,
        multi_spec_results=multi_spec_results,
        unserviceable_groups=unserviceable_groups,
    )


def execute_actual_production_workflow(
    adapted_data: AdaptedProductionData,
    routing_result: GroupRoutingResult,
    decoder_mode: str,
    output_dir: str | Path,
    panel_widths: list[int] | None = None,
    panel_width: int = PANEL_WIDTH,
    verbose: bool = True,
) -> WorkflowExecutionResult:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    runtime = load_runtime_modules()

    candidate_panel_widths = sorted(
        {
            int(width)
            for width in (panel_widths if panel_widths is not None else [panel_width])
            if int(width) > 0
        }
    )
    if not candidate_panel_widths:
        candidate_panel_widths = [int(panel_width)]

    apply_runtime_panel_width(runtime, candidate_panel_widths[0])
    apply_runtime_scale_factor(runtime, AUTO_LARGE_BATCH_SCALE_FACTOR)

    detail = adapted_data.printable_detail.copy()
    routing_summary = routing_result.routing_summary.copy()
    detail_group_id_col = detail.columns[0]

    route_group_id_col = routing_summary.columns[0]
    route_name_col = routing_summary.columns[1]
    route_thickness_col = routing_summary.columns[2]
    route_spec_count_col = routing_summary.columns[3]
    route_order_total_col = routing_summary.columns[4]
    route_weight_total_col = routing_summary.columns[5]
    route_route_type_col = routing_summary.columns[6]
    route_panel_width_col = routing_summary.columns[7]
    route_layout_note_col = routing_summary.columns[8]
    route_layout_width_col = routing_summary.columns[9]
    route_layout_length_col = routing_summary.columns[10]
    route_strip_lane_col = routing_summary.columns[11]
    route_strip_count_col = routing_summary.columns[12]
    route_output_total_col = routing_summary.columns[13]
    route_over_total_col = routing_summary.columns[14]
    route_consumed_length_col = routing_summary.columns[15]
    route_waste_width_col = routing_summary.columns[16]
    route_utilization_col = routing_summary.columns[17]
    route_note_col = routing_summary.columns[18]

    single_spec_mask = routing_summary[route_output_total_col].notna()
    single_spec_results = routing_summary.loc[
        single_spec_mask,
        [
            route_group_id_col,
            route_name_col,
            route_thickness_col,
            route_spec_count_col,
            route_order_total_col,
            route_weight_total_col,
            route_panel_width_col,
            route_layout_note_col,
            route_layout_width_col,
            route_layout_length_col,
            route_strip_lane_col,
            route_strip_count_col,
            route_output_total_col,
            route_over_total_col,
            route_consumed_length_col,
            route_waste_width_col,
            route_utilization_col,
            route_note_col,
        ],
    ].copy()
    if single_spec_results.empty:
        single_spec_results = pd.DataFrame(columns=SINGLE_SPEC_RESULT_COLUMNS)
    else:
        single_spec_results["订单数"] = single_spec_results[route_order_total_col].astype(int)
        single_spec_results["算法产出数"] = single_spec_results[route_output_total_col].astype(int)
        single_spec_results["补切数"] = (
            single_spec_results["订单数"] - single_spec_results["算法产出数"]
        ).clip(lower=0)
        single_spec_results["超产数"] = (
            single_spec_results["算法产出数"] - single_spec_results["订单数"]
        ).clip(lower=0)
        consumed_area = (
            single_spec_results[route_panel_width_col].fillna(0)
            * single_spec_results[route_consumed_length_col].fillna(0)
        )
        produced_area = (
            single_spec_results["算法产出数"].fillna(0)
            * single_spec_results[route_layout_width_col].fillna(0)
            * single_spec_results[route_layout_length_col].fillna(0)
        )
        single_spec_results["真实利用率(%)"] = 0.0
        valid_mask = consumed_area > 0
        single_spec_results.loc[valid_mask, "真实利用率(%)"] = (
            100.0 * produced_area.loc[valid_mask] / consumed_area.loc[valid_mask]
        )
        single_spec_results = single_spec_results.rename(
            columns={
                route_group_id_col: SS_GROUP_ID_COL,
                route_name_col: SINGLE_SPEC_RESULT_COLUMNS[1],
                route_thickness_col: SINGLE_SPEC_RESULT_COLUMNS[2],
                route_spec_count_col: SINGLE_SPEC_RESULT_COLUMNS[3],
                route_weight_total_col: SINGLE_SPEC_RESULT_COLUMNS[8],
                route_panel_width_col: SS_PANEL_WIDTH_COL,
                route_layout_note_col: SINGLE_SPEC_RESULT_COLUMNS[10],
                route_layout_width_col: SINGLE_SPEC_RESULT_COLUMNS[11],
                route_layout_length_col: SINGLE_SPEC_RESULT_COLUMNS[12],
                route_strip_lane_col: SINGLE_SPEC_RESULT_COLUMNS[13],
                route_strip_count_col: SINGLE_SPEC_RESULT_COLUMNS[14],
                route_consumed_length_col: SS_CONSUMED_LENGTH_COL,
                route_waste_width_col: SINGLE_SPEC_RESULT_COLUMNS[16],
                route_note_col: SINGLE_SPEC_RESULT_COLUMNS[18],
                "订单数": SINGLE_SPEC_RESULT_COLUMNS[4],
                "算法产出数": SINGLE_SPEC_RESULT_COLUMNS[5],
                "补切数": SINGLE_SPEC_RESULT_COLUMNS[6],
                "超产数": SINGLE_SPEC_RESULT_COLUMNS[7],
                "真实利用率(%)": SS_UTILIZATION_COL,
            }
        )
        single_spec_results = single_spec_results.loc[:, SINGLE_SPEC_RESULT_COLUMNS]

    multi_spec_rows: list[dict] = []
    workflow_rows: list[dict] = []

    for group_row in routing_summary.to_dict(orient="records"):
        group_id = group_row[route_group_id_col]

        if pd.notna(group_row[route_output_total_col]):
            single_spec_output_dir = output_dir / f"width_{int(group_row[route_panel_width_col])}"
            image_path, report_path = _save_single_spec_outputs(
                group_id=str(group_row[route_group_id_col]),
                name=str(group_row[route_name_col]),
                thickness=float(group_row[route_thickness_col]),
                panel_width=float(group_row[route_panel_width_col]),
                orientation=str(group_row[route_layout_note_col]),
                layout_width=float(group_row[route_layout_width_col]),
                layout_length=float(group_row[route_layout_length_col]),
                lane_count=int(group_row[route_strip_lane_col]),
                strip_count=int(group_row[route_strip_count_col]),
                order_total=int(group_row[route_order_total_col]),
                produced_total=int(group_row[route_output_total_col]),
                makeup_total=max(
                    int(group_row[route_order_total_col]) - int(group_row[route_output_total_col]),
                    0,
                ),
                over_total=max(
                    int(group_row[route_output_total_col]) - int(group_row[route_order_total_col]),
                    0,
                ),
                weight_total=float(group_row[route_weight_total_col]),
                consumed_length=float(group_row[route_consumed_length_col]),
                waste_width=float(group_row[route_waste_width_col]),
                utilization=float(group_row[route_utilization_col]),
                output_dir=single_spec_output_dir,
            )
            workflow_rows.append(
                {
                    WF_GROUP_ID_COL: group_row[route_group_id_col],
                    WF_NAME_COL: group_row[route_name_col],
                    WF_THICKNESS_COL: group_row[route_thickness_col],
                    WF_SPEC_COUNT_COL: group_row[route_spec_count_col],
                    WF_ROUTE_TYPE_COL: group_row[route_route_type_col],
                    WF_STATUS_COL: "完成",
                    WF_STRATEGY_BUCKET_COL: "单规格直接排版",
                    WF_DECODER_MODE_COL: "",
                    WF_SCALE_FACTOR_COL: None,
                    WF_PANEL_WIDTH_COL: group_row[route_panel_width_col],
                    WF_LAYOUT_NOTE_COL: group_row[route_layout_note_col],
                    WF_ORDER_TOTAL_COL: int(group_row[route_order_total_col]),
                    WF_PRODUCED_TOTAL_COL: int(group_row[route_output_total_col]),
                    WF_MAKEUP_TOTAL_COL: max(
                        int(group_row[route_order_total_col]) - int(group_row[route_output_total_col]),
                        0,
                    ),
                    WF_OVER_TOTAL_COL: max(
                        int(group_row[route_output_total_col]) - int(group_row[route_order_total_col]),
                        0,
                    ),
                    WF_WEIGHT_TOTAL_COL: group_row[route_weight_total_col],
                    WF_CONSUMED_LENGTH_COL: group_row[route_consumed_length_col],
                    WF_UTILIZATION_COL: (
                        100.0
                        * float(group_row[route_output_total_col])
                        * float(group_row[route_layout_width_col])
                        * float(group_row[route_layout_length_col])
                        / (
                            float(group_row[route_panel_width_col])
                            * float(group_row[route_consumed_length_col])
                        )
                        if group_row[route_panel_width_col] and group_row[route_consumed_length_col]
                        else 0.0
                    ),
                    WF_STRIP_COUNT_COL: group_row[route_strip_count_col],
                    WF_IMAGE_PATH_COL: str(image_path),
                    WF_REPORT_PATH_COL: str(report_path),
                    WF_NOTE_COL: group_row[route_note_col],
                }
            )
            continue

        group_detail = detail.loc[detail[detail_group_id_col] == group_id].copy()
        group_strategy = choose_multi_spec_strategy(
            group_row=group_row,
            requested_decoder_mode=decoder_mode,
        )

        if verbose:
            attempted_text = ", ".join(str(width) for width in candidate_panel_widths)
            print(
                f"[候选宽度比较] {group_id} "
                f"{group_row[route_name_col]} 厚度 {group_row[route_thickness_col]} "
                f"规格种数 {group_row[route_spec_count_col]} -> {attempted_text}"
            )

        successful_attempts: list[dict] = []
        failed_attempts: list[dict] = []
        for candidate_width in candidate_panel_widths:
            try:
                apply_runtime_panel_width(runtime, candidate_width)
                orientation_candidates = _build_multi_spec_orientation_candidates(
                    group_detail=group_detail,
                    panel_width=candidate_width,
                )
            except Exception as exc:
                failed_attempts.append(
                    {
                        "panel_width": candidate_width,
                        "orientation_note": "",
                        "error": str(exc),
                    }
                )
                if verbose:
                    print(f"  -> {candidate_width} mm 失败: {exc}")
                continue

            for orientation_candidate in orientation_candidates:
                try:
                    attempt_result = run_multi_spec_group(
                        group_row=group_row,
                        group_detail=group_detail,
                        decoder_mode=decoder_mode,
                        runtime=runtime,
                        output_dir=output_dir,
                        panel_width=candidate_width,
                        save_outputs=False,
                        orientation_choices=orientation_candidate["orientation_choices"],
                        orientation_note=orientation_candidate["orientation_note"],
                    )
                    successful_attempts.append(
                        {
                            "panel_width": candidate_width,
                            "orientation_note": orientation_candidate["orientation_note"],
                            "rotation_enabled": orientation_candidate["rotation_enabled"],
                            "result": attempt_result,
                        }
                    )
                    if verbose:
                        print(
                            f"  -> {candidate_width} mm / {orientation_candidate['orientation_note']} 成功: "
                            f"补切 {attempt_result[MS_MAKEUP_TOTAL_COL]} "
                            f"超产 {attempt_result[MS_OVER_TOTAL_COL]} 利用率 {attempt_result[MS_UTILIZATION_COL]:.4f}%"
                        )
                except Exception as exc:
                    failed_attempts.append(
                        {
                            "panel_width": candidate_width,
                            "orientation_note": orientation_candidate["orientation_note"],
                            "error": str(exc),
                        }
                    )
                    if verbose:
                        print(
                            f"  -> {candidate_width} mm / {orientation_candidate['orientation_note']} 失败: {exc}"
                        )

        if successful_attempts:
            best_attempt = min(successful_attempts, key=lambda item: _score_multi_spec_result(item["result"]))
            multi_result = best_attempt["result"]
            best_output_dir = output_dir / f"width_{best_attempt['panel_width']}"
            best_output_dir.mkdir(parents=True, exist_ok=True)
            save_context = multi_result.get("__save_context__", {})
            image_path, report_path = _save_multi_spec_outputs(
                runtime=runtime,
                group_id=save_context["group_id"],
                decoder_mode=save_context["decoder_mode"],
                demand=save_context["demand"],
                result={
                    "solution": save_context["solution"],
                    "real_efficiency": save_context["real_efficiency"],
                },
                output_dir=best_output_dir,
                panel_width=best_attempt["panel_width"],
            )
            multi_result[MS_IMAGE_PATH_COL] = str(image_path)
            multi_result[MS_REPORT_PATH_COL] = str(report_path)
            attempted_text = ", ".join(str(width) for width in candidate_panel_widths)
            rotation_text = (
                "启用，仅规格数=2时枚举可行朝向"
                if any(attempt.get("rotation_enabled") for attempt in successful_attempts)
                else "未启用或无可行旋转候选"
            )
            multi_result[MS_NOTE_COL] = (
                f"{multi_result[MS_NOTE_COL]}；候选宽度比较：{attempted_text}；"
                f"选定母板宽度：{best_attempt['panel_width']} mm；"
                f"旋转候选：{rotation_text}；最终朝向：{best_attempt['orientation_note']}"
            )
            multi_spec_rows.append(multi_result)
            workflow_rows.append(
                {
                    WF_GROUP_ID_COL: multi_result[MS_GROUP_ID_COL],
                    WF_NAME_COL: multi_result[MS_NAME_COL],
                    WF_THICKNESS_COL: multi_result[MS_THICKNESS_COL],
                    WF_SPEC_COUNT_COL: multi_result[MS_SPEC_COUNT_COL],
                    WF_ROUTE_TYPE_COL: group_row[route_route_type_col],
                    WF_STATUS_COL: "完成",
                    WF_STRATEGY_BUCKET_COL: multi_result[MS_STRATEGY_BUCKET_COL],
                    WF_DECODER_MODE_COL: multi_result[MS_DECODER_MODE_COL],
                    WF_SCALE_FACTOR_COL: multi_result[MS_SCALE_FACTOR_COL],
                    WF_PANEL_WIDTH_COL: multi_result[MS_PANEL_WIDTH_COL],
                    WF_LAYOUT_NOTE_COL: best_attempt["orientation_note"],
                    WF_ORDER_TOTAL_COL: multi_result[MS_ORDER_TOTAL_COL],
                    WF_PRODUCED_TOTAL_COL: multi_result[MS_PRODUCED_TOTAL_COL],
                    WF_MAKEUP_TOTAL_COL: multi_result[MS_MAKEUP_TOTAL_COL],
                    WF_OVER_TOTAL_COL: multi_result[MS_OVER_TOTAL_COL],
                    WF_WEIGHT_TOTAL_COL: multi_result[MS_WEIGHT_TOTAL_COL],
                    WF_CONSUMED_LENGTH_COL: multi_result[MS_REAL_CONSUMED_LENGTH_COL],
                    WF_UTILIZATION_COL: multi_result[MS_UTILIZATION_COL],
                    WF_STRIP_COUNT_COL: multi_result[MS_STRIP_COUNT_COL],
                    WF_IMAGE_PATH_COL: multi_result[MS_IMAGE_PATH_COL],
                    WF_REPORT_PATH_COL: multi_result[MS_REPORT_PATH_COL],
                    WF_NOTE_COL: multi_result[MS_NOTE_COL],
                }
            )
            continue

        attempted_text = ", ".join(str(width) for width in candidate_panel_widths)
        failure_note = failed_attempts[-1]["error"] if failed_attempts else "候选母板宽度均未生成可用方案"
        workflow_rows.append(
            {
                WF_GROUP_ID_COL: group_row[route_group_id_col],
                WF_NAME_COL: group_row[route_name_col],
                WF_THICKNESS_COL: group_row[route_thickness_col],
                WF_SPEC_COUNT_COL: group_row[route_spec_count_col],
                WF_ROUTE_TYPE_COL: group_row[route_route_type_col],
                WF_STATUS_COL: "失败",
                WF_STRATEGY_BUCKET_COL: group_strategy["strategy_bucket"],
                WF_DECODER_MODE_COL: group_strategy["decoder_mode"],
                WF_SCALE_FACTOR_COL: group_strategy["scale_factor"],
                WF_PANEL_WIDTH_COL: max(candidate_panel_widths),
                WF_LAYOUT_NOTE_COL: "",
                WF_ORDER_TOTAL_COL: group_row[route_order_total_col],
                WF_PRODUCED_TOTAL_COL: None,
                WF_MAKEUP_TOTAL_COL: None,
                WF_OVER_TOTAL_COL: None,
                WF_WEIGHT_TOTAL_COL: group_row[route_weight_total_col],
                WF_CONSUMED_LENGTH_COL: None,
                WF_UTILIZATION_COL: None,
                WF_STRIP_COUNT_COL: None,
                WF_IMAGE_PATH_COL: "",
                WF_REPORT_PATH_COL: "",
                WF_NOTE_COL: f"{failure_note}；候选宽度：{attempted_text}",
            }
        )

    workflow_summary = pd.DataFrame(workflow_rows)
    if workflow_summary.empty:
        workflow_summary = pd.DataFrame(columns=WORKFLOW_SUMMARY_COLUMNS)
    else:
        workflow_summary = attach_business_decisions(workflow_summary)
        workflow_summary = workflow_summary.loc[:, WORKFLOW_SUMMARY_COLUMNS].sort_values(
            [WF_GROUP_ID_COL],
            ignore_index=True,
        )

    multi_spec_results = pd.DataFrame(multi_spec_rows)
    if multi_spec_results.empty:
        multi_spec_results = pd.DataFrame(columns=MULTI_SPEC_RESULT_COLUMNS)
    else:
        multi_spec_results = multi_spec_results.loc[:, MULTI_SPEC_RESULT_COLUMNS].sort_values(
            [MS_GROUP_ID_COL],
            ignore_index=True,
        )

    width_allocation_summary = _build_width_allocation_summary(workflow_summary)
    unserviceable_groups = _build_unserviceable_groups(
        adapted_data=adapted_data,
        workflow_summary=workflow_summary,
    )

    return WorkflowExecutionResult(
        workflow_summary=workflow_summary,
        single_spec_results=single_spec_results,
        multi_spec_results=multi_spec_results,
        unserviceable_groups=unserviceable_groups,
        width_allocation_summary=width_allocation_summary,
    )


def export_actual_production_workflow(
    adapted_data: AdaptedProductionData,
    routing_result: GroupRoutingResult,
    execution_result: WorkflowExecutionResult,
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        adapted_data.printable_detail.to_excel(writer, sheet_name="可排版明细", index=False)
        adapted_data.excluded_detail.to_excel(writer, sheet_name="剔除明细", index=False)
        adapted_data.grouped_summary.to_excel(writer, sheet_name="按品名厚度分组结果", index=False)
        routing_result.routing_summary.to_excel(writer, sheet_name="分组处理策略", index=False)
        execution_result.workflow_summary.to_excel(writer, sheet_name="主流程汇总", index=False)
        execution_result.single_spec_results.to_excel(writer, sheet_name="单规格直排结果", index=False)
        execution_result.multi_spec_results.to_excel(writer, sheet_name="多规格GA结果", index=False)
        execution_result.unserviceable_groups.to_excel(writer, sheet_name="现有母板无法承接", index=False)

    return output_path


def export_actual_production_workflow(
    adapted_data: AdaptedProductionData,
    routing_result: GroupRoutingResult,
    execution_result: WorkflowExecutionResult,
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        adapted_data.printable_detail.to_excel(writer, sheet_name="可排版明细", index=False)
        adapted_data.excluded_detail.to_excel(writer, sheet_name="剔除明细", index=False)
        adapted_data.grouped_summary.to_excel(writer, sheet_name="按品名厚度分组", index=False)
        routing_result.routing_summary.to_excel(writer, sheet_name="分组路由", index=False)
        execution_result.workflow_summary.to_excel(writer, sheet_name="主流程汇总", index=False)
        execution_result.width_allocation_summary.to_excel(writer, sheet_name="宽度分配汇总", index=False)
        execution_result.single_spec_results.to_excel(writer, sheet_name="单规格结果", index=False)
        execution_result.multi_spec_results.to_excel(writer, sheet_name="多规格结果", index=False)
        execution_result.unserviceable_groups.to_excel(writer, sheet_name="无法承接分组", index=False)

    return output_path


def summarize_workflow_run(
    adapted_data: AdaptedProductionData,
    execution_result: WorkflowExecutionResult,
    panel_width: int,
    decoder_mode: str,
) -> WorkflowRunSummary:
    total_groups = len(execution_result.workflow_summary)
    completed_groups = int((execution_result.workflow_summary["状态"] == "已完成").sum())
    failed_groups = int((execution_result.workflow_summary["状态"] == "失败").sum())
    single_spec_groups = int((execution_result.workflow_summary["处理方式"] == "单规格直排").sum())
    multi_spec_groups = int((execution_result.workflow_summary["处理方式"] == "多规格优化").sum())

    total_demand_area = float(
        (
            adapted_data.printable_detail["短边"]
            * adapted_data.printable_detail["长边"]
            * adapted_data.printable_detail["片数"]
        ).sum()
    )

    total_consumed_length_mm = 0.0
    total_output_area = 0.0
    total_makeup_pieces = int(execution_result.workflow_summary["补切数"].fillna(0).sum())
    total_overproduction_pieces = int(execution_result.workflow_summary["超产数"].fillna(0).sum())
    if not execution_result.single_spec_results.empty:
        total_consumed_length_mm += float(execution_result.single_spec_results["总消耗长度(mm)"].fillna(0).sum())
        total_output_area += float(
            (
                execution_result.single_spec_results["算法产出总片数"].fillna(0)
                * execution_result.single_spec_results["最优横向占宽"].fillna(0)
                * execution_result.single_spec_results["最优纵向定尺"].fillna(0)
            ).sum()
        )
    if not execution_result.multi_spec_results.empty:
        total_consumed_length_mm += float(execution_result.multi_spec_results["还原后总长度(mm)"].fillna(0).sum())
        total_output_area += float(
            (
                execution_result.multi_spec_results["真实利用率(%)"].fillna(0)
                / 100.0
                * float(panel_width)
                * execution_result.multi_spec_results["还原后总长度(mm)"].fillna(0)
            ).sum()
        )

    total_panel_area = float(panel_width) * total_consumed_length_mm
    overall_efficiency = 0.0
    if total_panel_area > 0:
        overall_efficiency = 100.0 * total_output_area / total_panel_area

    return WorkflowRunSummary(
        panel_width=int(panel_width),
        decoder_mode=decoder_mode,
        total_groups=total_groups,
        completed_groups=completed_groups,
        failed_groups=failed_groups,
        single_spec_groups=single_spec_groups,
        multi_spec_groups=multi_spec_groups,
        total_demand_area=total_demand_area,
        total_output_area=total_output_area,
        total_makeup_pieces=total_makeup_pieces,
        total_overproduction_pieces=total_overproduction_pieces,
        total_consumed_length_mm=total_consumed_length_mm,
        total_panel_area=total_panel_area,
        overall_efficiency=overall_efficiency,
    )


def summarize_workflow_run(
    adapted_data: AdaptedProductionData,
    execution_result: WorkflowExecutionResult,
    panel_width: int,
    decoder_mode: str,
) -> WorkflowRunSummary:
    workflow_summary = execution_result.workflow_summary.copy()
    total_groups = len(workflow_summary)
    completed_groups = int(workflow_summary[WF_PRODUCED_TOTAL_COL].notna().sum())
    failed_groups = int(workflow_summary[WF_PRODUCED_TOTAL_COL].isna().sum())
    single_spec_groups = int(workflow_summary[WF_DECODER_MODE_COL].fillna("").astype(str).eq("").sum())
    multi_spec_groups = int(total_groups - single_spec_groups)

    total_makeup_pieces = int(workflow_summary[WF_MAKEUP_TOTAL_COL].fillna(0).sum())
    total_overproduction_pieces = int(workflow_summary[WF_OVER_TOTAL_COL].fillna(0).sum())
    total_consumed_length_mm = float(workflow_summary[WF_CONSUMED_LENGTH_COL].fillna(0).sum())

    panel_area_series = (
        workflow_summary[WF_PANEL_WIDTH_COL].fillna(0).astype(float)
        * workflow_summary[WF_CONSUMED_LENGTH_COL].fillna(0).astype(float)
    )
    output_area_series = (
        panel_area_series
        * workflow_summary[WF_UTILIZATION_COL].fillna(0).astype(float)
        / 100.0
    )
    total_panel_area = float(panel_area_series.sum())
    total_output_area = float(output_area_series.sum())
    overall_efficiency = 0.0
    if total_panel_area > 0:
        overall_efficiency = 100.0 * total_output_area / total_panel_area

    distinct_widths = sorted(
        {
            int(width)
            for width in workflow_summary[WF_PANEL_WIDTH_COL].dropna().tolist()
            if float(width) > 0
        }
    )
    summary_panel_width = int(distinct_widths[0]) if len(distinct_widths) == 1 else 0

    return WorkflowRunSummary(
        panel_width=summary_panel_width,
        decoder_mode=decoder_mode,
        total_groups=total_groups,
        completed_groups=completed_groups,
        failed_groups=failed_groups,
        single_spec_groups=single_spec_groups,
        multi_spec_groups=multi_spec_groups,
        total_demand_area=0.0,
        total_output_area=total_output_area,
        total_makeup_pieces=total_makeup_pieces,
        total_overproduction_pieces=total_overproduction_pieces,
        total_consumed_length_mm=total_consumed_length_mm,
        total_panel_area=total_panel_area,
        overall_efficiency=overall_efficiency,
    )


def summarize_workflow_run(
    adapted_data: AdaptedProductionData,
    execution_result: WorkflowExecutionResult,
    panel_width: int,
    decoder_mode: str,
) -> WorkflowRunSummary:
    workflow_summary = execution_result.workflow_summary.copy()
    total_groups = len(workflow_summary)
    completed_groups = int(workflow_summary[WF_PRODUCED_TOTAL_COL].notna().sum())
    failed_groups = int(workflow_summary[WF_PRODUCED_TOTAL_COL].isna().sum())
    single_spec_groups = int(workflow_summary[WF_DECODER_MODE_COL].fillna("").astype(str).eq("").sum())
    multi_spec_groups = int(total_groups - single_spec_groups)

    total_demand_area = 0.0

    total_makeup_pieces = int(workflow_summary[WF_MAKEUP_TOTAL_COL].fillna(0).sum())
    total_overproduction_pieces = int(workflow_summary[WF_OVER_TOTAL_COL].fillna(0).sum())
    total_consumed_length_mm = float(workflow_summary[WF_CONSUMED_LENGTH_COL].fillna(0).sum())

    panel_area_series = (
        workflow_summary[WF_PANEL_WIDTH_COL].fillna(0).astype(float)
        * workflow_summary[WF_CONSUMED_LENGTH_COL].fillna(0).astype(float)
    )
    output_area_series = (
        panel_area_series
        * workflow_summary[WF_UTILIZATION_COL].fillna(0).astype(float)
        / 100.0
    )
    total_panel_area = float(panel_area_series.sum())
    total_output_area = float(output_area_series.sum())
    overall_efficiency = 0.0
    if total_panel_area > 0:
        overall_efficiency = 100.0 * total_output_area / total_panel_area

    distinct_widths = sorted(
        {
            int(width)
            for width in workflow_summary[WF_PANEL_WIDTH_COL].dropna().tolist()
            if float(width) > 0
        }
    )
    summary_panel_width = int(distinct_widths[0]) if len(distinct_widths) == 1 else 0

    return WorkflowRunSummary(
        panel_width=summary_panel_width,
        decoder_mode=decoder_mode,
        total_groups=total_groups,
        completed_groups=completed_groups,
        failed_groups=failed_groups,
        single_spec_groups=single_spec_groups,
        multi_spec_groups=multi_spec_groups,
        total_demand_area=total_demand_area,
        total_output_area=total_output_area,
        total_makeup_pieces=total_makeup_pieces,
        total_overproduction_pieces=total_overproduction_pieces,
        total_consumed_length_mm=total_consumed_length_mm,
        total_panel_area=total_panel_area,
        overall_efficiency=overall_efficiency,
    )
