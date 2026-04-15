# -*- coding: utf-8 -*-
"""
Generate a client-facing PPT report for the GA steel cutting project.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Inches, Pt


ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from core.config import OUTPUT_DIR


PPT_FILE = OUTPUT_DIR / "ga_client_report.pptx"
CUTTING_PLAN_IMG = OUTPUT_DIR / "ga_fast_cutting_plan.png"
EVOLUTION_IMG = OUTPUT_DIR / "ga_fast_evolution_history.png"
DETAIL_IMG = OUTPUT_DIR / "ga_fast_strip_details.png"

DEMAND_ROWS = [
    ("T1", 230, 445, 1000),
    ("T2", 500, 833, 1000),
    ("T3", 547, 291, 500),
    ("T4", 200, 555, 666),
    ("T5", 400, 600, 888),
]

GA_RESULT = {
    "decoder_mode": "stage_based",
    "panel_width": 1250,
    "min_cut_gap": 200,
    "scale_factor": 3,
    "population_size": 24,
    "max_generations": 60,
    "crossover_rate": 0.85,
    "mutation_rate": 0.22,
    "elite_size": 3,
    "local_search_interval": 30,
    "early_stopping_patience": 35,
    "encoded_items": 1350,
    "best_fitness": 245024,
    "stop_generation": 41,
    "runtime_seconds": 0.98,
    "total_length_mm": 735072,
    "num_stages": 15,
    "efficiency_pct": 96.37,
    "demand_area": 885_484_500,
    "used_area": 918_840_000,
    "waste_area": 33_355_500,
    "mode_count": 5,
}

BASELINE_RESULT = {
    "greedy_length_mm": 817455,
    "greedy_efficiency_pct": 86.66,
    "reference_length_mm": 771390,
    "reference_efficiency_pct": 91.83,
}

PATTERN_ROWS = [
    ("模式 1", "T5×3", 59400, "96.00%"),
    ("模式 2", "T1×1 + T2×2", 139111, "98.40%"),
    ("模式 3", "T3×1 + T4×3", 41070, "91.76%"),
    ("模式 4", "T1×3 + T3×1", 3115, "98.96%"),
    ("模式 5", "T3×2", 2328, "87.52%"),
]

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

COLOR_NAVY = RGBColor(18, 42, 66)
COLOR_BLUE = RGBColor(35, 87, 137)
COLOR_TEAL = RGBColor(46, 125, 154)
COLOR_GOLD = RGBColor(201, 157, 60)
COLOR_RED = RGBColor(172, 47, 47)
COLOR_BG = RGBColor(245, 248, 252)
COLOR_TEXT = RGBColor(36, 44, 53)
COLOR_MUTED = RGBColor(92, 101, 112)
COLOR_LINE = RGBColor(210, 219, 230)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_GREEN = RGBColor(35, 123, 83)

FONT_CN = "Microsoft YaHei"


def mm_fmt(value: int | float) -> str:
    return f"{value:,.0f} mm"


def pct_fmt(value: float) -> str:
    return f"{value:.2f}%"


def add_full_background(slide, color=COLOR_BG):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_header_bar(slide, title: str, subtitle: str | None = None):
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, Cm(1.25)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_NAVY
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Cm(0.9), Cm(0.18), Cm(17), Cm(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = FONT_CN
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLOR_WHITE

    if subtitle:
        sub_box = slide.shapes.add_textbox(Cm(18.2), Cm(0.26), Cm(14), Cm(0.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = subtitle
        run.font.name = FONT_CN
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(214, 225, 238)


def add_bullets(slide, x, y, w, h, title, bullets, font_size=16, title_size=22):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    tf.margin_top = Cm(0.05)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = FONT_CN
    run.font.size = Pt(title_size)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY

    for bullet in bullets:
        p = tf.add_paragraph()
        p.level = 0
        p.space_before = Pt(6)
        p.bullet = True
        run = p.add_run()
        run.text = bullet
        run.font.name = FONT_CN
        run.font.size = Pt(font_size)
        run.font.color.rgb = COLOR_TEXT
    return box


def add_note_box(slide, x, y, w, h, title, body, fill=COLOR_WHITE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = COLOR_LINE

    box = slide.shapes.add_textbox(x + Cm(0.45), y + Cm(0.35), w - Cm(0.8), h - Cm(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = FONT_CN
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLUE

    for line in body:
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.name = FONT_CN
        run.font.size = Pt(11.5)
        run.font.color.rgb = COLOR_TEXT


def add_metric_cards(slide, cards):
    card_w = Cm(6.0)
    gap = Cm(0.45)
    start_x = Cm(0.9)
    y = Cm(1.7)
    for idx, (label, value, sub, color) in enumerate(cards):
        x = start_x + idx * (card_w + gap)
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, card_w, Cm(2.55)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_LINE

        top = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, card_w, Cm(0.22)
        )
        top.fill.solid()
        top.fill.fore_color.rgb = color
        top.line.fill.background()

        box = slide.shapes.add_textbox(x + Cm(0.4), y + Cm(0.35), card_w - Cm(0.7), Cm(1.9))
        tf = box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = label
        run.font.name = FONT_CN
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_MUTED

        p = tf.add_paragraph()
        run = p.add_run()
        run.text = value
        run.font.name = FONT_CN
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY

        p = tf.add_paragraph()
        p.space_before = Pt(1)
        run = p.add_run()
        run.text = sub
        run.font.name = FONT_CN
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_MUTED


def add_table(slide, x, y, w, h, headers, rows, col_widths=None):
    rows_count = len(rows) + 1
    cols_count = len(headers)
    table = slide.shapes.add_table(rows_count, cols_count, x, y, w, h).table

    if col_widths:
        for idx, col_w in enumerate(col_widths):
            table.columns[idx].width = col_w

    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = FONT_CN
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = COLOR_WHITE
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE if row_idx % 2 else RGBColor(248, 250, 253)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_CN
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = COLOR_TEXT
            cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return table


def add_image_or_placeholder(slide, image_path: Path, x, y, w, h, title: str):
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), x, y, w, h)
    else:
        placeholder = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h
        )
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(250, 252, 255)
        placeholder.line.color.rgb = COLOR_LINE

        box = slide.shapes.add_textbox(x + Cm(0.5), y + Cm(0.7), w - Cm(1.0), h - Cm(1.2))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"{title}\n\n图片未找到：{image_path.name}"
        run.font.name = FONT_CN
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_MUTED


def add_compare_bars(slide, x, y, w, labels_and_values, max_value):
    row_h = Cm(1.2)
    for idx, (label, value, color, suffix) in enumerate(labels_and_values):
        row_y = y + idx * row_h
        label_box = slide.shapes.add_textbox(x, row_y + Cm(0.15), Cm(3.4), Cm(0.6))
        tf = label_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = label
        run.font.name = FONT_CN
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_TEXT

        track_x = x + Cm(3.5)
        track = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, track_x, row_y + Cm(0.1), w, Cm(0.55)
        )
        track.fill.solid()
        track.fill.fore_color.rgb = RGBColor(232, 238, 246)
        track.line.fill.background()

        bar_w = w * (value / max_value)
        bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, track_x, row_y + Cm(0.1), bar_w, Cm(0.55)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        val_box = slide.shapes.add_textbox(track_x + w + Cm(0.2), row_y + Cm(0.04), Cm(3.7), Cm(0.7))
        tf = val_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = f"{value:,.2f}{suffix}"
        run.font.name = FONT_CN
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY


def add_section_slide(
    prs: Presentation,
    title: str,
    subtitle: str,
    left_title: str,
    left_bullets,
    right_title: str,
    right_bullets,
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, title, subtitle)
    add_bullets(slide, Cm(1.0), Cm(1.8), Cm(14.7), Cm(13.5), left_title, left_bullets)
    add_bullets(slide, Cm(16.5), Cm(1.8), Cm(15.4), Cm(13.5), right_title, right_bullets)
    return slide


def add_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, COLOR_BG)

    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Cm(1), SLIDE_H
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_GOLD
    accent.line.fill.background()

    navy_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Cm(1.4), Cm(0.9), Cm(15.5), Cm(15.8)
    )
    navy_panel.fill.solid()
    navy_panel.fill.fore_color.rgb = COLOR_NAVY
    navy_panel.line.fill.background()

    title_box = slide.shapes.add_textbox(Cm(2.4), Cm(2.1), Cm(13), Cm(5.2))
    tf = title_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "基于遗传算法的\n钢板组刀与纵切优化项目汇报"
    run.font.name = FONT_CN
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = COLOR_WHITE

    p2 = tf.add_paragraph()
    p2.space_before = Pt(12)
    run = p2.add_run()
    run.text = "面向甲方的项目背景、算法设计、实验结果与实施建议"
    run.font.name = FONT_CN
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(222, 231, 240)

    info_box = slide.shapes.add_textbox(Cm(2.4), Cm(10.1), Cm(12.5), Cm(2.8))
    tf = info_box.text_frame
    for idx, line in enumerate(
        [
            "应用场景：多规格订单钢板切割 / 组刀优化",
            f"当前方案：快速 GA + {GA_RESULT['decoder_mode']} 解码",
            "汇报定位：方案可落地性、效果价值、实施路径",
        ]
    ):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = FONT_CN
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_WHITE

    summary_panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Cm(18.0), Cm(1.2), Cm(14.0), Cm(13.8)
    )
    summary_panel.fill.solid()
    summary_panel.fill.fore_color.rgb = COLOR_WHITE
    summary_panel.line.color.rgb = COLOR_LINE

    box = slide.shapes.add_textbox(Cm(19.0), Cm(2.0), Cm(12.0), Cm(11.5))
    tf = box.text_frame
    lines = [
        "项目一句话说明",
        "在满足工艺约束与订单需求的前提下，自动生成材料利用率更高、可执行性更强的钢板切割方案。",
        "",
        "本次样例核心结果",
        f"总切割长度：{mm_fmt(GA_RESULT['total_length_mm'])}",
        f"材料利用率：{pct_fmt(GA_RESULT['efficiency_pct'])}",
        f"阶段模板数量：{GA_RESULT['mode_count']} 类",
        f"求解时间：{GA_RESULT['runtime_seconds']:.2f} s",
    ]
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        run.font.name = FONT_CN
        run.font.size = Pt(14 if idx in (0, 3) else 12.5)
        run.font.bold = idx in (0, 3)
        run.font.color.rgb = COLOR_NAVY if idx in (0, 3) else COLOR_TEXT


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    add_title_slide(prs)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "执行摘要", "Management Summary")
    add_metric_cards(
        slide,
        [
            ("总切割长度", mm_fmt(GA_RESULT["total_length_mm"]), "实际用料长度越低越优", COLOR_BLUE),
            ("材料利用率", pct_fmt(GA_RESULT["efficiency_pct"]), "需求面积 / 实际使用面积", COLOR_GREEN),
            ("阶段模板数量", str(GA_RESULT["mode_count"]), "方案结构已较好收敛", COLOR_GOLD),
            ("求解时间", f"{GA_RESULT['runtime_seconds']:.2f} s", "支持快速出方案", COLOR_TEAL),
            ("停止代数", f"第 {GA_RESULT['stop_generation']} 代", "触发早停，搜索趋于稳定", COLOR_RED),
        ],
    )
    add_note_box(
        slide,
        Cm(1.0),
        Cm(4.8),
        Cm(14.8),
        Cm(7.8),
        "核心结论",
        [
            "当前项目已经形成一套可直接服务实际生产决策的 GA 组刀优化流程，而不是停留在算法验证层面。",
            "本方案将遗传算法搜索与阶段式解码器结合，使输出结果更接近生产现场的连续切割逻辑。",
            "在当前样例中，方案材料利用率达到 96.37%，且求解时间控制在 1 秒左右，具备工程实用性。",
        ],
    )
    add_note_box(
        slide,
        Cm(16.2),
        Cm(4.8),
        Cm(15.0),
        Cm(7.8),
        "面向甲方的价值",
        [
            "降低钢板消耗和浪费面积，直接带来原材料成本改善。",
            "降低人工组刀对经验的依赖，提升方案稳定性与可复盘性。",
            "为后续与 ERP / MES 对接、形成标准模式库提供基础。",
        ],
        fill=RGBColor(250, 252, 255),
    )

    add_section_slide(
        prs,
        "项目背景",
        "Business Background",
        "业务场景",
        [
            "钢板切割业务同时面对多规格、多长度、多批量订单并存的生产现实，人工组刀难以兼顾全局最优。",
            "传统经验排刀往往更关注单次可切性，而难以系统平衡宽度填充、长度损耗、尾单收口和阶段切换。",
            "在订单波动频繁、交期紧张时，人工方案既耗时，也容易出现不同班组之间结果不一致的问题。",
        ],
        "主要痛点",
        [
            "排刀效率不足：规格一多，人工试排需要反复试错。",
            "材料浪费不可控：单看宽度填充率并不代表整体用料最优。",
            "方案稳定性不足：不同人员、不同时间的经验方案差异大。",
            "缺少可追溯性：很难沉淀出一套标准模式与评价依据。",
        ],
    )

    add_section_slide(
        prs,
        "问题定义",
        "Problem Statement",
        "拟解决的问题",
        [
            "在满足所有订单需求的前提下，自动给出钢板切割与组刀方案，目标是最小化总切割长度和材料浪费。",
            "方案不仅要可行，还要尽量贴近现场工艺执行方式，避免输出只适合论文、不适合生产的结果。",
            "最终输出应当包含阶段切割组合、阶段长度、模式明细和可视化结果，方便工艺与生产人员复核。",
        ],
        "关键工艺约束",
        [
            f"大板宽度固定为 {GA_RESULT['panel_width']} mm。",
            f"最小纵切约束为 {GA_RESULT['min_cut_gap']} mm，过窄组合将被判为违规并施加高额罚项。",
            "不同产品长度不一致时，会引入长度方向损耗，因此不能只追求宽度拼满。",
            "需要精确满足不同产品的需求数量，同时尽量减少阶段切换与尾单模式。",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "输入数据与业务样例", "Input Scenario")
    add_bullets(
        slide,
        Cm(1.0),
        Cm(1.8),
        Cm(10.5),
        Cm(4.8),
        "样例数据说明",
        [
            "数据来源为 Excel 订单表，字段包括 Width、Length、num、Weight。",
            "当前样例共 5 类产品，适合展示多规格并行订单下的切割优化逻辑。",
            "项目对需求采用缩放编码，当前缩放因子为 3，以兼顾求解速度与结果稳定性。",
        ],
        font_size=14,
    )
    add_table(
        slide,
        Cm(1.0),
        Cm(6.0),
        Cm(15.0),
        Cm(7.6),
        ["产品", "宽度(mm)", "长度(mm)", "需求数量"],
        DEMAND_ROWS,
        [Cm(2.2), Cm(3.0), Cm(3.0), Cm(4.2)],
    )
    add_note_box(
        slide,
        Cm(17.0),
        Cm(2.1),
        Cm(14.2),
        Cm(11.2),
        "当前数据规模",
        [
            f"编码后小板数量：{GA_RESULT['encoded_items']}",
            f"需求总面积：{GA_RESULT['demand_area']:,} mm²",
            "订单组合同时包含短料、中料和长料，长度差异明显，因此非常适合作为算法验证场景。",
            "这类数据特点决定了：如果只按宽度贪心，通常会带来明显的长度浪费。",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "数学建模", "Optimization Model")
    add_note_box(
        slide,
        Cm(1.0),
        Cm(1.8),
        Cm(14.8),
        Cm(5.2),
        "目标函数",
        [
            "可将问题抽象为：min TotalLength + Penalty。",
            "其中 TotalLength 表示所有阶段切割长度之和，Penalty 用于处罚违反工艺约束的组合。",
            "适应度越小表示方案越优，因此算法会持续朝着低总长度、低违规风险方向进化。",
        ],
        fill=RGBColor(255, 253, 248),
    )
    add_note_box(
        slide,
        Cm(1.0),
        Cm(7.4),
        Cm(14.8),
        Cm(5.6),
        "主要约束",
        [
            f"宽度约束：任一阶段宽度组合之和不得超过 {GA_RESULT['panel_width']} mm。",
            f"最小纵切约束：参与纵切的产品宽度需满足不小于 {GA_RESULT['min_cut_gap']} mm。",
            "需求约束：各产品累计产出数量需覆盖订单需求。",
            "工艺合理性约束：长度差异较大的混排会引入更高的长度浪费。",
        ],
    )
    add_bullets(
        slide,
        Cm(17.0),
        Cm(2.0),
        Cm(14.0),
        Cm(11.5),
        "为什么不用简单穷举",
        [
            "该问题本质上是离散组合优化，随着产品数量和需求规模上升，搜索空间呈指数级膨胀。",
            "如果用穷举或纯规则法，很难在可接受时间内同时兼顾全局用料和局部工艺结构。",
            "因此项目采用遗传算法搜索与专用解码器结合的方式，在工程上更具可行性。",
        ],
        font_size=14,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "算法总体路线", "GA Workflow")
    add_note_box(
        slide,
        Cm(1.0),
        Cm(1.8),
        Cm(30.2),
        Cm(3.0),
        "总体流程",
        [
            "Excel需求 -> 数据展开/缩放 -> 染色体编码 -> 解码生成阶段方案 -> 适应度评估 -> 选择/交叉/变异 -> 局部搜索 -> 输出切割图与明细",
        ],
        fill=RGBColor(248, 252, 255),
    )
    stages = [
        ("1", "需求建模", "读取订单宽度、长度、数量等信息"),
        ("2", "染色体编码", "用产品排列顺序表示一个候选解"),
        ("3", "解码评估", "将排列映射为阶段式切割方案"),
        ("4", "进化迭代", "通过选择、交叉、变异不断改进"),
        ("5", "局部搜索", "对当前最优解做轻量精修"),
        ("6", "结果输出", "生成图表、模板和汇报材料"),
    ]
    start_x = Cm(1.0)
    y = Cm(6.0)
    box_w = Cm(4.85)
    gap = Cm(0.25)
    for idx, (num, title, desc) in enumerate(stages):
        x = start_x + idx * (box_w + gap)
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, box_w, Cm(5.8)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_LINE

        badge = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, x + Cm(0.35), y + Cm(0.35), Cm(0.9), Cm(0.9)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_BLUE
        badge.line.fill.background()
        t = slide.shapes.add_textbox(x + Cm(0.35), y + Cm(0.42), Cm(0.9), Cm(0.5))
        p = t.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = num
        run.font.name = FONT_CN
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_WHITE

        box = slide.shapes.add_textbox(x + Cm(0.45), y + Cm(1.5), box_w - Cm(0.8), Cm(3.8))
        tf = box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.name = FONT_CN
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY

        p = tf.add_paragraph()
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.name = FONT_CN
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_TEXT

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "染色体编码与遗传操作", "Encoding and Operators")
    add_bullets(
        slide,
        Cm(1.0),
        Cm(1.8),
        Cm(14.7),
        Cm(12.5),
        "编码设计",
        [
            "染色体采用产品索引排列编码，一个个体就是所有待切小板的一个排列顺序。",
            "这种编码方式便于交叉与变异，同时将复杂工艺规则交由解码器处理，结构清晰、扩展性好。",
            "适应度通过解码后的总切割长度与罚项共同决定，越小越优。",
        ],
        font_size=14,
    )
    add_bullets(
        slide,
        Cm(16.5),
        Cm(1.8),
        Cm(15.0),
        Cm(12.5),
        "遗传操作",
        [
            "选择：采用锦标赛选择，增强优秀个体被继承的概率。",
            "交叉：采用 OX 顺序交叉，保留排列相对顺序，适合排列编码问题。",
            "变异：采用 swap / insert / inversion 混合变异，提高跳出局部最优能力。",
            "精英保留：每代保留最优个体，避免优秀方案退化。",
        ],
        font_size=14,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "项目核心创新：阶段式解码器", "Stage-based Decoder")
    add_note_box(
        slide,
        Cm(1.0),
        Cm(1.8),
        Cm(14.6),
        Cm(10.8),
        "为什么它是项目亮点",
        [
            "当前实现并不是把个体简单地翻译成一条条静态 strip，而是翻译成阶段式连续生产方案。",
            "算法会先根据染色体顺序生成当前阶段的宽度组合，再计算该组合可以连续切割的阶段长度。",
            "当某一类产品在该阶段先被满足后，再自动切换到下一阶段组合。",
            "这让输出结果更接近现场的工艺执行逻辑，也更容易被甲方接受为可落地方案。",
        ],
        fill=RGBColor(250, 252, 255),
    )
    add_note_box(
        slide,
        Cm(16.1),
        Cm(1.8),
        Cm(15.1),
        Cm(10.8),
        "对业务的实际意义",
        [
            "减少单次排刀结果与实际生产组织之间的脱节。",
            "更容易形成阶段模板加阶段长度的标准化工艺表达方式。",
            "有利于后续与现场作业指导单、MES 下发逻辑或标准模式库对接。",
            "相比只优化单条宽度填充率，这种方式更能兼顾整体用料与执行便利性。",
        ],
        fill=RGBColor(255, 253, 248),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "工程化优化设计", "Engineering Enhancements")
    add_metric_cards(
        slide,
        [
            ("种群规模", str(GA_RESULT["population_size"]), "控制计算量与解质量平衡", COLOR_BLUE),
            ("最大迭代", str(GA_RESULT["max_generations"]), "允许充分进化", COLOR_TEAL),
            ("交叉率", f"{GA_RESULT['crossover_rate']:.2f}", "强化优秀结构重组", COLOR_GOLD),
            ("变异率", f"{GA_RESULT['mutation_rate']:.2f}", "维持种群多样性", COLOR_RED),
            ("精英保留", str(GA_RESULT["elite_size"]), "防止优解退化", COLOR_GREEN),
        ],
    )
    add_bullets(
        slide,
        Cm(1.0),
        Cm(4.6),
        Cm(15.0),
        Cm(8.0),
        "增强机制",
        [
            "贪心初始化：按宽度、长度、类型等规则生成种子个体，提升初始解质量。",
            "局部搜索：每隔 30 代对当前最优解进行 swap 与 block move 微调。",
            "动态变异率：在长期不改进时自动提高变异概率，增强跳出局部最优能力。",
            "早停机制：连续 35 代无改进时停止，避免无效计算。",
        ],
        font_size=14,
    )
    add_note_box(
        slide,
        Cm(16.5),
        Cm(4.8),
        Cm(14.6),
        Cm(7.6),
        "工程结论",
        [
            "这套实现并不是原始教科书版 GA，而是面向大规模实例做过求解速度和稳定性增强的工程版本。",
            "在当前样例下，算法不到 1 秒即可完成搜索，说明具备日常使用的潜力。",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "实验设置", "Experiment Setup")
    add_bullets(
        slide,
        Cm(1.0),
        Cm(1.9),
        Cm(14.6),
        Cm(11.8),
        "实验环境与设置",
        [
            "数据：5 类产品、多长度混合需求。",
            f"解码模式：{GA_RESULT['decoder_mode']}。",
            f"缩放因子：{GA_RESULT['scale_factor']}，编码后样本量 {GA_RESULT['encoded_items']}。",
            "随机种子固定为 42，确保结果可复现。",
            "快速版 GA 输出进化曲线图、阶段排版图和条带明细图，用于工程验证和汇报展示。",
        ],
        font_size=14,
    )
    add_note_box(
        slide,
        Cm(16.4),
        Cm(2.0),
        Cm(14.8),
        Cm(5.1),
        "评估指标",
        [
            "总切割长度：越低越优。",
            "材料利用率：越高越优。",
            "模式数量：越少通常越利于生产执行与标准化。",
            "求解时间：越短越利于现场快速响应。",
        ],
    )
    add_note_box(
        slide,
        Cm(16.4),
        Cm(7.5),
        Cm(14.8),
        Cm(4.8),
        "判定原则",
        [
            "本项目以可执行性、材料效率和计算速度三方面综合判断方案优劣。",
            "因此不会只拿单一宽度利用率做评价，而是看整体用料长度和阶段结构。",
        ],
        fill=RGBColor(255, 253, 248),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "实验结果总览", "Results Overview")
    add_metric_cards(
        slide,
        [
            ("最优适应度", f"{GA_RESULT['best_fitness']:,}", "Fitness = TotalLength + Penalty", COLOR_BLUE),
            ("总切割长度", mm_fmt(GA_RESULT["total_length_mm"]), "还原后真实用料长度", COLOR_GREEN),
            ("材料利用率", pct_fmt(GA_RESULT["efficiency_pct"]), "需求面积 / 实际使用面积", COLOR_GOLD),
            ("浪费面积", f"{GA_RESULT['waste_area']:,} mm²", "已压降到较低水平", COLOR_RED),
            ("模式数量", str(GA_RESULT["mode_count"]), "输出结构较清晰", COLOR_TEAL),
        ],
    )
    add_note_box(
        slide,
        Cm(1.0),
        Cm(4.8),
        Cm(14.8),
        Cm(7.6),
        "结果解读",
        [
            f"算法在第 {GA_RESULT['stop_generation']} 代触发早停，说明搜索已较稳定。",
            f"总切割长度控制在 {GA_RESULT['total_length_mm']:,} mm，对应材料利用率 {GA_RESULT['efficiency_pct']:.2f}%。",
            f"实际使用面积为 {GA_RESULT['used_area']:,} mm²，其中浪费面积为 {GA_RESULT['waste_area']:,} mm²。",
            "从业务视角看，这组结果体现了较好的材料效率和较清晰的生产阶段结构。",
        ],
    )
    add_note_box(
        slide,
        Cm(16.2),
        Cm(4.8),
        Cm(15.0),
        Cm(7.6),
        "与项目内参考方案对比",
        [
            f"相对贪婪参考方案，总切割长度减少 {BASELINE_RESULT['greedy_length_mm'] - GA_RESULT['total_length_mm']:,} mm。",
            f"相对宽度最大化参考方案，总切割长度减少 {BASELINE_RESULT['reference_length_mm'] - GA_RESULT['total_length_mm']:,} mm。",
            "说明 GA 方案不仅在局部宽度上做得好，更在全局阶段组织上取得优势。",
        ],
        fill=RGBColor(250, 252, 255),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "方案对比分析", "Comparison")
    add_note_box(
        slide,
        Cm(1.0),
        Cm(1.8),
        Cm(14.6),
        Cm(11.5),
        "对比结论",
        [
            "对比对象 1：项目内贪婪参考方案，特点是规则简单、速度快，但容易忽略全局长度损耗。",
            "对比对象 2：仓库中的宽度最大化参考方案，能够得到较好的局部组合，但阶段组织能力有限。",
            "当前 GA 方案在总切割长度与利用率两项核心指标上均表现更优。",
        ],
    )
    add_compare_bars(
        slide,
        Cm(17.0),
        Cm(2.2),
        Cm(8.4),
        [
            ("GA 方案长度", GA_RESULT["total_length_mm"], COLOR_BLUE, " mm"),
            ("贪婪方案长度", BASELINE_RESULT["greedy_length_mm"], COLOR_RED, " mm"),
            ("参考方案长度", BASELINE_RESULT["reference_length_mm"], COLOR_GOLD, " mm"),
        ],
        max(
            GA_RESULT["total_length_mm"],
            BASELINE_RESULT["greedy_length_mm"],
            BASELINE_RESULT["reference_length_mm"],
        ),
    )
    add_compare_bars(
        slide,
        Cm(17.0),
        Cm(6.4),
        Cm(8.4),
        [
            ("GA 利用率", GA_RESULT["efficiency_pct"], COLOR_GREEN, "%"),
            ("贪婪利用率", BASELINE_RESULT["greedy_efficiency_pct"], COLOR_RED, "%"),
            ("参考利用率", BASELINE_RESULT["reference_efficiency_pct"], COLOR_GOLD, "%"),
        ],
        100.0,
    )
    add_note_box(
        slide,
        Cm(17.0),
        Cm(10.2),
        Cm(14.2),
        Cm(3.1),
        "甲方可直接理解的价值表达",
        [
            "同样的需求量下，当前方案用更短的总切割长度完成生产，本质上就是更省料、更稳、更易标准化。",
        ],
        fill=RGBColor(255, 253, 248),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "结果可视化：阶段排版图", "Cutting Plan Visual")
    add_image_or_placeholder(slide, CUTTING_PLAN_IMG, Cm(1.0), Cm(1.8), Cm(19.5), Cm(11.6), "阶段排版图")
    add_note_box(
        slide,
        Cm(21.0),
        Cm(2.0),
        Cm(10.2),
        Cm(11.2),
        "图示解读",
        [
            "左侧图片展示的是阶段式排版结果，而非单条静态 strip 的简单堆叠。",
            "每一个阶段对应一组宽度组合和一段连续生产长度，能够直接映射到现场执行层。",
            "这类图对于甲方非常重要，因为它直观体现了算法输出是可生产的，而不仅是一个抽象评分。",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "结果可视化：进化过程", "Evolution History")
    add_image_or_placeholder(slide, EVOLUTION_IMG, Cm(1.0), Cm(2.0), Cm(18.8), Cm(10.6), "进化曲线图")
    add_note_box(
        slide,
        Cm(20.3),
        Cm(2.0),
        Cm(11.0),
        Cm(10.6),
        "进化过程说明",
        [
            "从进化曲线可以看到，算法在前期快速下降，随后趋于平稳。",
            f"在第 {GA_RESULT['stop_generation']} 代左右已基本收敛，因此触发早停，避免无效迭代。",
            "这说明当前参数设置在收敛质量和运行速度之间达到了较好的平衡。",
        ],
        fill=RGBColor(250, 252, 255),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "阶段模板明细", "Pattern Summary")
    add_table(
        slide,
        Cm(1.0),
        Cm(2.0),
        Cm(19.0),
        Cm(9.2),
        ["模板", "组成", "阶段长度(mm)", "宽度利用率"],
        PATTERN_ROWS,
        [Cm(2.5), Cm(7.6), Cm(4.5), Cm(4.0)],
    )
    add_image_or_placeholder(slide, DETAIL_IMG, Cm(20.6), Cm(2.0), Cm(10.4), Cm(9.2), "条带明细图")
    add_note_box(
        slide,
        Cm(1.0),
        Cm(11.6),
        Cm(30.0),
        Cm(1.7),
        "模式结构结论",
        [
            "主方案优先采用高利用率模板承担主要产量，低利用率模板只用于尾单收口与需求平衡，这种结构非常符合工程常识。",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide)
    add_header_bar(slide, "实施建议", "Deployment Recommendations")
    add_bullets(
        slide,
        Cm(1.0),
        Cm(1.9),
        Cm(14.8),
        Cm(11.6),
        "落地建议",
        [
            "建议先将本方案作为智能排刀辅助决策系统上线，而不是一开始就完全替代人工。",
            "推荐采用算法自动出方案加工艺人员审核确认的协同模式，以降低上线阻力。",
            "日常订单使用快速 GA 即可；超大规模订单可提高缩放因子后先粗排，再对重点订单精修。",
            "建议沉淀高频优质模式库，逐步形成标准模板，进一步减少现场沟通成本。",
        ],
        font_size=14,
    )
    add_bullets(
        slide,
        Cm(16.5),
        Cm(1.9),
        Cm(15.0),
        Cm(11.6),
        "下一阶段工作",
        [
            "对接 ERP / MES，自动读取订单并自动输出标准排刀单。",
            "引入更多现场约束，例如换刀成本、批次切换成本、设备节拍等。",
            "开展 2 到 4 周的现场 A/B 测试，用真实订单验证节材率和排刀效率提升。",
            "建立算法方案、实际执行和生产反馈的闭环，不断迭代模型。",
        ],
        font_size=14,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_background(slide, RGBColor(247, 250, 253))
    add_header_bar(slide, "汇报结论", "Closing")
    add_note_box(
        slide,
        Cm(2.4),
        Cm(2.0),
        Cm(27.2),
        Cm(8.6),
        "总结陈述",
        [
            "本项目的核心价值，不是单纯把遗传算法引入工厂，而是把复杂、多规格、动态变化的组刀问题，转化成一套可自动求解、可解释、可执行的生产优化流程。",
            f"在当前样例中，系统以约 {GA_RESULT['runtime_seconds']:.2f} 秒的时间给出材料利用率 {GA_RESULT['efficiency_pct']:.2f}% 的阶段式方案，证明其具备算法先进性与工程可用性。",
            "因此，该方案适合作为甲方后续建设智能排刀、智能排产和工艺标准化能力的基础模块。",
        ],
        fill=COLOR_WHITE,
    )
    end_box = slide.shapes.add_textbox(Cm(2.4), Cm(11.3), Cm(26), Cm(2.0))
    tf = end_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "谢谢。可继续展开为现场试运行方案、ROI测算和系统集成方案。"
    run.font.name = FONT_CN
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY

    return prs


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(PPT_FILE)
    print(f"PPT 已生成：{PPT_FILE}")


if __name__ == "__main__":
    main()
