# -*- coding: utf-8 -*-
"""
运行指定 sheet 与指定解码模式的 GA 实验，并输出最终切割排布图。

示例：
python scripts/run_decoder_experiments.py --sheet 4 --modes simple best_fit
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from core.config import (  # noqa: E402
    CROSSOVER_RATE,
    DEFAULT_DATA_FILE,
    ELITE_SIZE,
    MAX_GENERATIONS,
    MUTATION_RATE,
    OUTPUT_DIR,
    PANEL_WIDTH,
    POPULATION_SIZE,
    SCALE_FACTOR,
)
from core.cutting_report import write_cutting_report  # noqa: E402
from core.data_loader import expand_demand  # noqa: E402
from core.decoder import merge_same_pattern_strips  # noqa: E402
from core.ga_engine_fast import GeneticAlgorithm  # noqa: E402
from core.visualization import (  # noqa: E402
    plot_compact_cutting_plan,
    plot_stage_based_cutting_plan,
)


def load_sheet_demand(sheet_num):
    """
    从 Excel 读取指定 sheet。
    某些 sheet 的重量列为空；当前 GA 目标函数不依赖重量，因此这里兜底补为 1.0。
    """
    wb = load_workbook(DEFAULT_DATA_FILE, data_only=True)
    sheet_name = f"Sheet{sheet_num}"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"未找到工作表 {sheet_name}")

    ws = wb[sheet_name]
    item_num = ws["B1"].value
    if item_num is None or item_num <= 0:
        raise ValueError(f"{sheet_name} 的产品数量读取失败")

    demand = {"Width": [], "Length": [], "num": [], "Weight": []}
    for i in range(item_num):
        row = i + 3
        width = ws[f"A{row}"].value
        length = ws[f"B{row}"].value
        num = ws[f"C{row}"].value
        weight = ws[f"D{row}"].value
        if width is None or length is None or num is None:
            raise ValueError(f"{sheet_name} 第 {i + 1} 行关键字段缺失")

        demand["Width"].append(width)
        demand["Length"].append(length)
        demand["num"].append(num)
        demand["Weight"].append(weight if weight not in (None, 0) else 1.0)

    return demand


def run_single_mode(sheet_num, items, demand, decoder_mode):
    ga = GeneticAlgorithm(
        items,
        population_size=POPULATION_SIZE,
        max_generations=MAX_GENERATIONS,
        crossover_rate=CROSSOVER_RATE,
        mutation_rate=MUTATION_RATE,
        elite_size=ELITE_SIZE,
        decoder_mode=decoder_mode,
    )
    ga.evolve(verbose=True)
    solution = ga.get_solution()

    real_total_length = solution["total_length"] * SCALE_FACTOR
    total_area_demand = sum(
        width * length * num
        for width, length, num in zip(demand["Width"], demand["Length"], demand["num"])
    )
    total_area_used = PANEL_WIDTH * real_total_length
    real_efficiency = 100 * total_area_demand / total_area_used if total_area_used else 0.0

    output_path = OUTPUT_DIR / f"sheet{sheet_num}_{decoder_mode}_cutting_plan.png"
    report_path = OUTPUT_DIR / f"sheet{sheet_num}_{decoder_mode}_cutting_report.txt"
    if decoder_mode == "stage_based":
        plot_stage_based_cutting_plan(
            solution["strips"],
            real_efficiency,
            save_path=output_path,
            show=False,
        )
    else:
        merged_strips, repeat_counts = merge_same_pattern_strips(solution["strips"])
        plot_compact_cutting_plan(
            merged_strips,
            repeat_counts,
            real_efficiency,
            save_path=output_path,
            show=False,
        )

    write_cutting_report(
        sheet_num=sheet_num,
        decoder_mode=decoder_mode,
        demand=demand,
        solution=solution,
        output_path=report_path,
    )

    return {
        "decoder_mode": decoder_mode,
        "total_length": real_total_length,
        "efficiency": real_efficiency,
        "num_strips": solution["num_strips"] * SCALE_FACTOR,
        "output_path": output_path,
        "report_path": report_path,
    }


def parse_args():
    parser = argparse.ArgumentParser(description="Run GA decoder experiments for a given sheet.")
    parser.add_argument("--sheet", type=int, required=True, help="Excel sheet number, e.g. 4")
    parser.add_argument(
        "--modes",
        nargs="+",
        required=True,
        help="Decoder modes to run, e.g. simple best_fit stage_based",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    demand = load_sheet_demand(args.sheet)
    items, _ = expand_demand(demand)

    print("=" * 72)
    print(f"Sheet{args.sheet} 解码模式实验")
    print("=" * 72)
    print(f"缩放因子: {SCALE_FACTOR}")
    print(f"编码后小板数: {len(items)}")

    results = []
    for mode in args.modes:
        print("\n" + "-" * 72)
        print(f"运行解码模式: {mode}")
        print("-" * 72)
        result = run_single_mode(args.sheet, items, demand, mode)
        results.append(result)
        print(
            f"{mode} 完成: 总长度={result['total_length']:.0f} mm, "
            f"利用率={result['efficiency']:.2f}%, "
            f"图片={result['output_path']}, 报告={result['report_path']}"
        )

    print("\n" + "=" * 72)
    print("实验结果汇总")
    print("=" * 72)
    for result in results:
        print(
            f"{result['decoder_mode']:<12} "
            f"总长度={result['total_length']:.0f} mm, "
            f"利用率={result['efficiency']:.2f}%, "
            f"切割条数={result['num_strips']}, "
            f"图片={result['output_path'].name}, "
            f"报告={result['report_path'].name}"
        )


if __name__ == "__main__":
    main()
